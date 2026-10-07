"""Authored exports on licensed development originals, separate from blind scoring."""

import argparse
import io
import json
from decimal import Decimal
from pathlib import Path
from zipfile import ZipFile

import pypdfium2 as pdfium
from docx import Document
from evaluate_native_v3 import digest, same_place
from pypdf import PdfReader

from paperdelta import __version__
from paperdelta.analysis import check_project
from paperdelta.annotations import annotation_bytes, preview_annotations
from paperdelta.config import config_text
from paperdelta.docx_document import DocxDocument
from paperdelta.i18n import language_context
from paperdelta.models import Config, ReviewedTableIdentity
from paperdelta.pdf_document import PdfDocument
from paperdelta.storage import Project, json_text, sha256
from paperdelta.tables import reviewed_table_anchor

ROOT = Path(__file__).resolve().parents[1]
CORPUS = ROOT / "validation/native-v3"


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", required=True)
    args = parser.parse_args()
    output = (ROOT / args.out).resolve()
    assert output.is_relative_to(ROOT / "build") and not output.exists()
    output.mkdir(parents=True)
    gold = json.loads((CORPUS / "development-gold.json").read_text("utf-8"))["targets"]
    font = json.loads((CORPUS / "development-font-boxes.json").read_text("utf-8"))["boxes"]
    records = []
    for name in ("dev-001", "dev-074"):
        target = next(t for t in gold if t["id"] == name)
        kind = target["format"]
        source = CORPUS / "papers" / target["case"] / target["file"]
        raw = source.read_bytes()
        assert digest(source) == target["source_sha256"]
        document = (DocxDocument if kind == "docx" else PdfDocument)(target["file"], raw)
        expected = {**target["locator"], **({"bbox": font[name]} if kind == "pdf" else {})}
        candidates = [s for s in document.numbers() if same_place(s.locator, expected)]
        assert len(candidates) == 1 and candidates[0].text == target["text"]
        span = candidates[0]
        # This is an explicit authored identity review, not automatic mapping.
        # Original table: Iteration 6 / Variant rs2240026 / Marginal P-value.
        anchor = (
            reviewed_table_anchor(
                document, span, ReviewedTableIdentity(header_rows=2, row_prefix=["6", "rs2240026"])
            )
            if kind == "docx"
            else document.anchor_for_span(span)
        )
        project = Project(output / kind)
        project.write(target["file"], raw)
        project.write("results.csv", f"id,value\nselected,{target['value']}\n".encode())
        config = Config.model_validate(
            {
                "schema_version": 10,
                "paper": {"entry": target["file"]},
                "sources": {
                    "data": {
                        "path": "results.csv",
                        "format": "csv",
                        "primary_key": ["id"],
                        "columns": {"id": "string", "value": "decimal"},
                    }
                },
                "metrics": {
                    "selected": {
                        "source": "data",
                        "field": "value",
                        "where": {"id": "selected"},
                        "unit": "scalar",
                    }
                },
                "occurrences": {
                    "selected": {
                        "file": target["file"],
                        "metric": "selected",
                        "anchor": anchor.model_dump(),
                        "display": {
                            "kind": "scientific" if kind == "docx" else "decimal",
                            "places": 1 if kind == "docx" else 0,
                            "percent_symbol": False,
                        },
                    }
                },
            }
        )
        project.write("paperdelta.yaml", config_text(config).encode())
        before = check_project(project.root)
        assert before["occurrences"]["selected"]["status"] == "pass"
        changed = Decimal(target["value"]) + Decimal("0.01" if kind == "docx" else "1")
        project.write("results.csv", f"id,value\nselected,{changed}\n".encode())
        after = check_project(project.root)
        assert after["occurrences"]["selected"]["status"] == "mismatch"
        for language in ("en", "zh-CN"):
            with language_context(language):
                plan = preview_annotations(project, {"occurrences": ["selected"]})
            assert not plan["refused"] and len(plan["entries"]) == 1
            archive = annotation_bytes(project, plan)
            project.write(language + ".zip", archive, exclusive=True)
            with ZipFile(io.BytesIO(archive)) as packed:
                copied = packed.read(plan["copies"][0]["path"])
                manifest = json.loads(packed.read("review.json"))
            assert manifest["copies_sha256"][plan["copies"][0]["path"]] == sha256(copied)
            project.write(language + ".review." + kind, copied, exclusive=True)
            if kind == "pdf":
                original, reviewed = PdfReader(io.BytesIO(raw)), PdfReader(io.BytesIO(copied))
                assert len(original.pages) == len(reviewed.pages)
                for old, new in zip(original.pages, reviewed.pages, strict=True):
                    assert old.get_contents().get_data() == new.get_contents().get_data()
                    assert old.mediabox == new.mediabox and old.cropbox == new.cropbox
                with pdfium.PdfDocument(copied) as pdf:
                    pdf[0].render(scale=2, draw_annots=True).to_pil().save(
                        project.root / (language + ".png")
                    )
            else:
                old, new = Document(io.BytesIO(raw)), Document(io.BytesIO(copied))
                assert len(new.comments) == len(old.comments) + 1
                assert list(new.comments)[-1].text == plan["entries"][0]["note"]
                assert DocxDocument(source.name, copied).text == document.text
            assert project.read(target["file"]) == raw and digest(source) == target["source_sha256"]
            records.append(
                {
                    "format": kind,
                    "language": language,
                    "gold_id": name,
                    "source_sha256": target["source_sha256"],
                    "copy_sha256": sha256(copied),
                    "plan": plan,
                    "check": "pass",
                    "changed_evidence": "mismatch",
                    "original_unchanged": True,
                }
            )
    evidence = {
        "tool_version": __version__,
        "scope": "Four authored bilingual exports on two development originals. "
        "The Word identity is explicitly reviewed; this is separate from "
        "automatic-anchor study results.",
        "cases": records,
    }
    (output / "evidence.json").write_text(json_text(evidence), encoding="utf-8")
    print(json.dumps({"status": "passed", "cases": len(records)}))


if __name__ == "__main__":
    main()
