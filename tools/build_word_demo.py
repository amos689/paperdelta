"""Rebuild the original, deterministic Word demo bundled in the optional workflow."""

import io
import zipfile
from datetime import UTC, datetime
from pathlib import Path

from docx import Document

from paperdelta.builder import anchor_for_span, candidate_span
from paperdelta.config import config_text
from paperdelta.docx_document import DocxDocument
from paperdelta.models import Config, Display


def build():
    document = Document()
    document.core_properties.author = "amos689"
    document.core_properties.last_modified_by = "amos689"
    document.core_properties.comments = ""
    document.core_properties.title = "PaperDelta Word example"
    document.core_properties.created = datetime(2026, 10, 3, tzinfo=UTC)
    document.core_properties.modified = datetime(2026, 10, 3, tzinfo=UTC)
    document.add_heading("PaperDelta: evidence-linked Word review", 0)
    document.add_heading("Abstract", 1)
    paragraph = document.add_paragraph("Our method achieves ")
    paragraph.add_run("84.").bold = True
    paragraph.add_run("1%")
    paragraph.add_run(" accuracy on Data-A.")
    document.add_heading("Results", 1)
    document.add_paragraph("Our method outperforms Baseline on Data-A.")
    document.add_paragraph("Accuracy on the held-out test split (%)", style="Caption")
    table = document.add_table(rows=3, cols=2)
    table.style = "Light Shading Accent 1"
    for row, values in zip(
        table.rows, [("Method", "Accuracy"), ("Ours", "84.1"), ("Baseline", "81.0")], strict=True
    ):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    stream = io.BytesIO()
    document.save(stream)
    # ZIP timestamps otherwise change source hashes without changing the document.
    normalized = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(stream.getvalue())) as source,
        zipfile.ZipFile(normalized, "w") as output,
    ):
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, (2026, 10, 3, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            output.writestr(info, source.read(name))
    raw = normalized.getvalue()
    native = DocxDocument("paper.docx", raw)
    metrics = {
        "accuracy": {
            "source": "results",
            "field": "accuracy",
            "where": {"model": "Ours"},
            "unit": "fraction",
            "reduce": "mean",
            "expected_count": 3,
            "expected_seeds": [1, 2, 3],
        },
        "baseline": {
            "source": "results",
            "field": "accuracy",
            "where": {"model": "Baseline"},
            "unit": "fraction",
            "expected_count": 1,
        },
    }
    occurrences = {}
    for name, metric, number, percent_symbol in zip(
        ["abstract_accuracy", "table_accuracy", "table_baseline"],
        ["accuracy", "accuracy", "baseline"],
        native.numbers(),
        [True, False, False],
        strict=True,
    ):
        display = Display(kind="percent", percent_symbol=percent_symbol)
        span = candidate_span(native, number.to_dict(), display)
        occurrences[name] = {
            "file": "paper.docx",
            "metric": metric,
            "anchor": anchor_for_span(native, span).model_dump(),
            "display": display.model_dump(),
        }
    config = Config.model_validate(
        {
            "schema_version": 3,
            "paper": {"entry": "paper.docx"},
            "sources": {
                "results": {
                    "path": "results/metrics.csv",
                    "format": "csv",
                    "primary_key": ["model", "seed"],
                    "columns": {"model": "string", "seed": "integer", "accuracy": "decimal"},
                }
            },
            "metrics": metrics,
            "occurrences": occurrences,
            "claims": {
                "main_comparison": {
                    "file": "paper.docx",
                    "anchor": {"exact": "Our method outperforms Baseline on Data-A."},
                    "predicate": {"op": "greater_than", "left": "accuracy", "right": "baseline"},
                }
            },
        }
    )
    root = Path(__file__).resolve().parents[1] / "src/paperdelta/demo_word"
    (root / "results").mkdir(parents=True, exist_ok=True)
    (root / "paper.docx").write_bytes(raw)
    (root / "paperdelta.yaml").write_bytes(config_text(config).encode("utf-8"))
    (root / "results/metrics.csv").write_text(
        "model,seed,accuracy\nOurs,1,0.839\nOurs,2,0.841\nOurs,3,0.843\nBaseline,1,0.810\n",
        encoding="utf-8",
        newline="\n",
    )


if __name__ == "__main__":
    build()
