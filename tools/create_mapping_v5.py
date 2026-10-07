"""Create new owned mapping tasks before Agent tuning; never overwrite a suite.

Native reference positions use the recorded 1.9 parser on authored documents.
Experiment contracts and required abstentions come from the fixture specification,
not model outputs. These are synthetic tasks, not independently annotated papers.
Generation additionally requires openpyxl==3.1.5; checking Excel does not.
"""

from __future__ import annotations

import argparse
import csv
import io
import json
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path

from paperdelta import __version__, builder
from paperdelta.models import Display
from paperdelta.onboarding import inspect_proposal, scan_project
from paperdelta.storage import Project, json_text, sha256

FORMATS = {
    "development": [
        "latex",
        "markdown",
        "docx",
        "markdown",
        "quarto",
        "pdf",
        "docx",
        "latex",
        "quarto",
        "pdf",
        "markdown",
        "latex",
    ],
    "held-out": [
        "markdown",
        "latex",
        "docx",
        "quarto",
        "latex",
        "pdf",
        "markdown",
        "quarto",
        "docx",
        "latex",
        "pdf",
        "quarto",
    ],
}
FAMILIES = [
    "checkpoint_identity",
    "tsv_scalar_stale_display",
    "fraction_percent_table",
    "escaped_json_pointer",
    "percentage_point_difference",
    "pdf_raw_percentage",
    "excel_string_identity",
    "relative_percentage_change",
    "missing_required_seed",
    "missing_checkpoint_choice",
    "undeclared_source_unit",
    "missing_significance_evidence",
]
EXTENSIONS = {"latex": "tex", "markdown": "md", "quarto": "qmd", "docx": "docx", "pdf": "pdf"}


def save(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("xb") as stream:
        stream.write(value if isinstance(value, bytes) else value.encode("utf-8"))


def paper_bytes(format, prefix, literal, suffix, row_label, table):
    text = prefix + literal + suffix
    if format == "latex":
        return (text.replace("%", r"\%") + "\n").encode()
    if format in {"markdown", "quarto"}:
        return ("# Results\n\n" + text + "\n").encode()
    if format == "docx":
        from docx import Document

        document = Document()
        document.core_properties.author = "PaperDelta"
        document.core_properties.last_modified_by = "PaperDelta"
        document.core_properties.created = datetime(2026, 10, 7, tzinfo=UTC)
        document.core_properties.modified = datetime(2026, 10, 7, tzinfo=UTC)
        document.add_heading("Results", level=1)
        if table:
            document.add_paragraph(prefix.rstrip(": "))
            cells = document.add_table(rows=2, cols=2)
            cells.style = "Table Grid"
            cells.cell(0, 0).text = "Experiment"
            cells.cell(0, 1).text = "Score (%)"
            cells.cell(1, 0).text = row_label
            cells.cell(1, 1).text = literal
            document.add_paragraph(suffix.strip())
        else:
            document.add_paragraph(text)
        output = io.BytesIO()
        document.save(output)
        return output.getvalue()
    from reportlab.pdfgen.canvas import Canvas

    output = io.BytesIO()
    canvas = Canvas(output, invariant=True, pagesize=(900, 500))
    canvas.setTitle("Authored mapping evidence")
    canvas.setAuthor("PaperDelta")
    canvas.setFont("Helvetica", 12)
    canvas.drawString(40, 450, "Results")
    canvas.drawString(40, 410, text)
    canvas.showPage()
    canvas.save()
    return output.getvalue()


def source_files(case, family, split, names, values):
    dataset, model, checkpoint = names
    rows = []
    for variant, experiment, measurements in (
        (checkpoint, model, values),
        ("alternate", model, values),
        (checkpoint, "Control", ["0.650", "0.652", "0.654"]),
    ):
        for seed, value in zip(("001", "002", "003"), measurements, strict=True):
            if family == 9 and variant == checkpoint and experiment == model and seed == "003":
                continue
            rows.append(
                {
                    "dataset": dataset,
                    "model": experiment,
                    "checkpoint": variant,
                    "split": "test",
                    "seed": seed,
                    "score": value,
                }
            )
    if family == 6:
        for row in rows:
            row["score"] = str(Decimal(row["score"]) * 100)
    if family == 3:
        for row in rows:
            row["score_pct"] = str(Decimal(row["score"]) * 100)
    columns = {key: "string" for key in ("dataset", "model", "checkpoint", "split", "seed")}
    columns["score"] = "decimal"
    if family == 3:
        columns["score_pct"] = "decimal"
    keys = ["dataset", "model", "checkpoint", "split", "seed"]
    buffer = io.StringIO(newline="")
    writer = csv.DictWriter(
        buffer,
        fieldnames=list(columns),
        delimiter="\t" if family == 2 else ",",
        lineterminator="\n",
    )
    writer.writeheader()
    writer.writerows(rows)
    format = "tsv" if family == 2 else "csv"
    path = "results/measurements." + format
    save(case / "project" / path, buffer.getvalue())
    source = {"path": path, "format": format, "columns": columns, "primary_key": keys}
    if family == 4:
        path = "results/export.json"
        mean = sum(map(Decimal, values)) / 3
        save(
            case / "project" / path,
            json_text(
                {
                    "runs": {
                        "evaluation/test": {
                            "score.mean": mean,
                            "dataset": dataset,
                            "model": model,
                            "checkpoint": checkpoint,
                            "split": "test",
                            "seeds": ["001", "002", "003"],
                            "unit": "fraction",
                        }
                    }
                }
            ),
        )
        source = {"path": path, "format": "json", "columns": {}, "primary_key": []}
    if family == 7:
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "held results" if split == "held-out" else "development results"
        sheet.append(["Author export. Use the explicitly declared range below."])
        sheet.append(list(columns))
        for row in rows:
            sheet.append([row[key] for key in columns])
        workbook.properties.creator = "PaperDelta"
        output = io.BytesIO()
        workbook.save(output)
        path = "results/workbook.xlsx"
        save(case / "project" / path, output.getvalue())
        source.update(path=path, format="xlsx", sheet=sheet.title, cell_range="A2:F11")
    return source


def reference(project, source, family, names, literal, display, row_label):
    dataset, model, checkpoint = names
    choices = [("score", "percent" if family == 6 else "scalar" if family == 2 else "fraction")]
    if family == 3:
        choices.append(("score_pct", "percent"))
    references, controls = [], []
    for field, unit in choices:
        draft = builder.start_draft(project)
        draft = builder.add_source(project, draft, name="records", **source)
        where = {"dataset": dataset, "model": model, "checkpoint": checkpoint, "split": "test"}
        definition = {
            "source": "records",
            "field": field,
            "unit": unit,
            "reduce": "mean",
            "where": where,
            "expected_count": 3,
            "seed_column": "seed",
            "expected_seeds": ["001", "002", "003"],
        }
        if family == 4:
            definition.update(
                field="/runs/evaluation~1test/score.mean",
                reduce="unique",
                where={},
                expected_count=1,
                expected_seeds=None,
            )
        draft = builder.add_metric(project, draft, name="target", **definition)
        metric = "target"
        if family in {5, 8}:
            other = {**definition, "where": {**where, "model": "Control"}}
            draft = builder.add_metric(project, draft, name="control", **other)
            draft = builder.add_derived(
                project,
                draft,
                name="contrast",
                operation="percentage_point_difference"
                if family == 5
                else "relative_change_percent",
                left="target",
                right="control",
            )
            metric = "contrast"
        scan = scan_project(project)
        numeric = literal.removesuffix("%")
        selected = [
            item
            for item in scan["candidates"]
            if item["text"].replace(r"\%", "%").removesuffix("%") == numeric
        ]
        assert len(selected) == 1, (family, literal, [item["text"] for item in scan["candidates"]])
        candidate = selected[0]
        table_identity = {"header_rows": 1, "row_prefix": [row_label]} if family == 3 else None
        draft = builder.add_occurrences(
            project,
            draft,
            metric=metric,
            candidate_ids=[candidate["candidate_id"]],
            names=["requested"],
            display_kind=display.kind,
            places=display.places,
            percent_symbol=display.percent_symbol,
            rationale=(
                "Authored fixture reference: independently specified source identity, "
                "units and requested original location."
            ),
            table_identity=table_identity,
        )
        proposal = builder.finalize_draft(project, draft)
        _, _, checked = inspect_proposal(project, proposal)
        entry = checked["occurrences"]["requested"]
        assert entry["status"] in {"pass", "mismatch"}
        references.append({key: proposal[key] for key in ("additions", "rationale")})
        controls.append(
            {
                "status": entry["status"],
                "location": entry["location"],
                "expected": entry["expected"],
                "candidate_id": candidate["candidate_id"],
            }
        )
    return references, controls


def chinese_request(family, entry, dataset, model, topic):
    subjects = {
        "mean test accuracy": "测试准确率的均值",
        "mean test loss": "测试损失的均值",
        "test accuracy gain in percentage points": "测试准确率提升的百分点差",
        "relative test accuracy improvement": "测试准确率的相对百分比变化",
        "significance p-value": "显著性检验的 p 值",
    }
    lines = [
        f"只绑定 {entry} 中原来的{subjects[topic]}位置。",
        f"数据集为 {dataset}，模型为 {model}，评估划分为 test。"
        "必需种子是字符串 001、002、003，保留前导零。",
    ]
    if family != 10:
        lines.append("已复核选择的 checkpoint 为 selected。")
    if family in {5, 8}:
        lines.append(
            "将选定模型与同一数据集、checkpoint、划分和完整种子集合上的 Control 比较。"
            + (
                "报告选定模型减去 Control 的百分点差。"
                if family == 5
                else "相对变化为（选定模型减去 Control）/ Control，显示为百分数。"
            )
        )
    extra = {
        3: "使用 Word 表格的 Score (%) 列和完整 Experiment 行标签确定单元格。首行为唯一表头。",
        9: "导出结果缺少一个必需种子，不能悄悄缩小预期运行集合。",
        10: (
            "两个 checkpoint 都存在且数值相同，但作者未记录正文用了哪一个。不能用数值相同决定身份。"
        ),
        11: (
            "来源单位声明未保留。正文百分号不能说明原始 score 是 fraction 还是 percent，"
            "不能按量级猜。"
        ),
        12: (
            "仅保留聚合准确率。原始观测、假设、检验方法及其元数据缺失，不能编造或计算显著性结果。"
        ),
    }
    if family in extra:
        lines.append(extra[family])
    lines.append(
        "正文数值过期仍可能具有正确映射。若不能确认身份、完整运行集合、单位或所需推导，"
        "请明确拒答并解释缺少的证据。保留全部文件；不能自动接受建议。"
    )
    return "# 映射请求\n\n" + "\n\n".join(lines) + "\n"


def build_case(output, split, family):
    tag = "D" if split == "development" else "H"
    case_id = tag + f"{family:02d}"
    case = output / split / case_id
    format = FORMATS[split][family - 1]
    names = (
        "Cedar" if tag == "D" else "Juniper",
        "VariantBlue" if tag == "D" else "VariantAmber",
        "selected",
    )
    values = ["0.712", "0.714", "0.716"] if tag == "D" else ["0.681", "0.683", "0.685"]
    if family == 2:
        values = ["0.7121", "0.7141", "0.7161"] if tag == "D" else ["0.6812", "0.6832", "0.6852"]
    source = source_files(case, family, split, names, values)
    mean, baseline = sum(map(Decimal, values)) / 3, Decimal("0.652")
    display = Display(kind="percent", places=1, percent_symbol=True)
    literal = f"{mean * 100:.1f}%"
    topic = "mean test accuracy"
    if family == 2:
        display = Display(kind="decimal", places=4, percent_symbol=False)
        literal = "0.7000" if tag == "D" else "0.6900"
        topic = "mean test loss"
    elif family == 5:
        display = Display(kind="decimal", places=1, percent_symbol=False)
        literal = f"{(mean - baseline) * 100:.1f}"
        topic = "test accuracy gain in percentage points"
    elif family == 8:
        literal = f"{(mean - baseline) / baseline * 100:.1f}%"
        topic = "relative test accuracy improvement"
    elif family == 12:
        display = Display(kind="decimal", places=3, percent_symbol=False)
        literal, topic = "0.032", "significance p-value"
    dataset, model, checkpoint = names
    row_label = f"{dataset} / {model} / {checkpoint} / test"
    prefix = f"The {topic} is "
    suffix = f" on {dataset} for {model}."
    entry = "paper/manuscript." + EXTENSIONS[format]
    save(
        case / "project" / entry,
        paper_bytes(format, prefix, literal, suffix, row_label, family == 3),
    )
    save(
        case / "project/paperdelta.yaml",
        json_text(
            {
                "schema_version": 4
                if format == "pdf"
                else 3
                if format == "docx"
                else 10
                if format in {"markdown", "quarto"}
                else 1,
                "paper": {"entry": entry},
            }
        ),
    )
    instructions = [
        f"Bind only the original {topic} result in {entry}.",
        f"Dataset: {dataset}. Model: {model}. Evaluation split: test. "
        "Required seed IDs are the exact strings 001, 002 and 003.",
    ]
    if family != 10:
        instructions.append("The reviewed run choice is checkpoint 'selected'.")
    if family in {5, 8}:
        instructions.append(
            "Compare the chosen model with Control on the same dataset, checkpoint, "
            "split and complete seeds. "
            + (
                "Report percentage-point difference, chosen minus Control."
                if family == 5
                else "Report relative change: (chosen minus Control) / Control, "
                "expressed as a percentage."
            )
        )
    if family == 3:
        instructions.append(
            "Use the cell in the Word table's Score (%) column and the complete "
            "Experiment row label. The first row is the sole header row."
        )
    if family == 9:
        instructions.append(
            "The export is missing a required seed. Do not silently reduce the expected run set."
        )
    if family == 10:
        instructions.append(
            "Both checkpoints exist and have the same values. The author did not record "
            "which checkpoint the manuscript used. Do not choose a checkpoint by number equality."
        )
    if family == 11:
        instructions.append(
            "No source unit declaration was retained. The manuscript percentage does not "
            "determine whether raw scores are fractions or percentages. "
            "Do not infer units from magnitude."
        )
    if family == 12:
        instructions.append(
            "Only aggregate accuracy scores were retained. Original observations, hypothesis, "
            "statistical test and test metadata are absent. "
            "Do not manufacture or calculate a significance result."
        )
    instructions.append(
        "A stale manuscript value may still have a correct mapping. If identity, complete runs, "
        "units or required derivation cannot be established, explicitly abstain and explain "
        "the missing evidence. Preserve all files; do not accept a proposal."
    )
    language = "zh-CN" if (family % 2 == 0) == (tag == "D") else "en"
    request = "# Mapping request\n\n" + "\n\n".join(instructions) + "\n"
    if language == "zh-CN":
        request = chinese_request(family, entry, dataset, model, topic)
    save(case / "request.md", request)
    notes = [
        "# Author's data export notes",
        "Record identity columns are dataset, model, checkpoint, split, seed. "
        "All five are strings; preserve leading zeros in seed IDs. "
        "Measurement fields are decimal values.",
        f"The authoritative evidence file for the requested result is {source['path']}.",
    ]
    if family == 4:
        notes.append(
            "Use the saved mean at runs -> evaluation/test -> score.mean in the JSON export. "
            "The neighboring CSV is a separate trial log. JSON values are fractions; "
            "its explicit metadata identifies the experiment and complete seed set."
        )
    elif family == 7:
        notes.append(
            f"Select worksheet {source['sheet']!r}, range {source['cell_range']}. "
            "The first row of that range is the column header; "
            "cells below it are string identifiers and decimal measurements."
        )
    if family != 11:
        notes.append(
            "score stores "
            + (
                "scalar loss"
                if family == 2
                else "percent accuracy (0 to 100)"
                if family == 6
                else "fraction accuracy (0 to 1)"
            )
            + "."
        )
    if family == 3:
        notes.append(
            "score_pct stores equivalent percent accuracy. Either field is acceptable "
            "when its source unit and manuscript formatting are declared correctly."
        )
    if family == 11:
        notes.append("The unit of score is unknown. No other unit annotation is available.")
    save(case / "project/EXPERIMENT.md", "\n\n".join(notes) + "\n")
    project = Project(case / "project")
    references, controls = (
        reference(project, source, family, names, literal, display, row_label)
        if family <= 8
        else ([], [])
    )
    oracle = {
        "schema_version": 1,
        "case_id": case_id,
        "family": FAMILIES[family - 1],
        "expected_decision": "map" if family <= 8 else "abstain",
        "reference_inputs": references,
        "reference_controls": controls,
        "required_seed_ids": ["001", "002", "003"],
        "native_position_origin": (
            "Authored template and pre-2.0 PaperDelta 1.9 positions; "
            "not independent native-parser gold."
        ),
        "reason": instructions[-2]
        if family >= 9
        else "Match the declared experiment contract and requested original position; "
        "a stale display remains mappable.",
    }
    save(case / "oracle.json", json_text(oracle))
    return {
        "id": case_id,
        "split": split,
        "family": FAMILIES[family - 1],
        "manuscript_format": format,
        "request_language": language,
        "expected_decision": oracle["expected_decision"],
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    assert __version__ == "1.9.0", "Generate before changing the Agent implementation"
    assert not args.out.exists(), "Use a new directory; preserve every previous attempt"
    output = args.out.resolve()
    output.mkdir(parents=True)
    cases = [build_case(output, split, family) for split in FORMATS for family in range(1, 13)]
    hashes = {
        path.relative_to(output).as_posix(): sha256(path.read_bytes())
        for path in sorted(output.rglob("*"))
        if path.is_file()
    }
    import platform
    from importlib.metadata import version

    import paperdelta

    package = Path(paperdelta.__file__).parent
    runtime = {
        path.relative_to(package).as_posix(): sha256(path.read_bytes())
        for path in sorted(package.rglob("*"))
        if path.is_file() and path.suffix in {".py", ".js", ".json", ".css"}
    }
    save(
        output / "manifest.json",
        json_text(
            {
                "suite_id": "mapping-v5",
                "created_at": datetime.now(UTC).isoformat(),
                "generator_version": __version__,
                "generator_sha256": sha256(Path(__file__).read_bytes()),
                "python": platform.python_version(),
                "generation_dependencies": {
                    name: version(name)
                    for name in ("python-docx", "pdfplumber", "reportlab", "openpyxl")
                },
                "pre_tuning_runtime_sha256": runtime,
                "license": "MIT",
                "cases": cases,
                "input_sha256": hashes,
                "scope": (
                    "Owned synthetic development and held-out tasks, 8 mappable "
                    "and 4 required abstentions per split. No model calls. "
                    "Gold is excluded from inference context."
                ),
            }
        ),
    )
    print(json.dumps({"status": "created", "cases": len(cases), "files": len(hashes)}))


if __name__ == "__main__":
    main()
