"""Authored regressions for native positions; not general parser accuracy claims."""

import io

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen.canvas import Canvas

from paperdelta.docx_document import DocxDocument
from paperdelta.errors import PaperDeltaError
from paperdelta.models import ReviewedTableIdentity
from paperdelta.pdf_document import PdfDocument
from paperdelta.tables import reviewed_table_anchor


def word_table(values, *, before=0, after=0, declared=True):
    document = Document()
    document.add_paragraph("Table 1. Experiment results")
    table = document.add_table(rows=1, cols=3)
    for cell, text in zip(table.rows[0].cells, ("Experiment", "Summary", "Comment"), strict=True):
        cell.text = text
    for label, result in values:
        cells = table.add_row().cells
        cells[0].text, cells[1].text, cells[2].text = label, result, "observed"
    row = table.rows[0]._tr
    if before or after:
        cells = row.findall(qn("w:tc"))
        for cell in cells[:before] + (cells[-after:] if after else []):
            row.remove(cell)
        if declared:
            for count, side in ((before, "Before"), (after, "After")):
                if count:
                    prop = OxmlElement("w:grid" + side)
                    prop.set(qn("w:val"), str(count))
                    row.get_or_add_trPr().append(prop)
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


@pytest.mark.parametrize(("before", "after"), [(1, 0), (0, 1), (1, 1)])
def test_explicit_omitted_header_cells_keep_original_columns(before, after):
    raw = word_table([("Ours", "84.1"), ("Baseline", "80.9")], before=before, after=after)
    doc = DocxDocument("p.docx", raw)
    target = next(s for s in doc.numbers() if s.text == "84.1")
    assert target.locator["cell"] == 2
    anchor = doc.anchor_for_span(target)
    assert doc.locate(anchor) == target
    assert not any(i["code"] == "DOCX_COMPLEX_TABLE" for i in doc.issues)
    assert doc.raw == raw


def test_undeclared_grid_holes_remain_unknown():
    doc = DocxDocument("p.docx", word_table([("Ours", "84.1")], before=1, declared=False))
    assert all(s.text != "84.1" for s in doc.numbers())
    assert any(i["code"] == "DOCX_COMPLEX_TABLE" for i in doc.issues)


def test_mean_spread_and_n_keep_independent_positions_when_all_values_change():
    doc = DocxDocument("p.docx", word_table([("Ours", "84.1 ± 1.2 (n = 3)")]))
    original = [s for s in doc.numbers() if s.locator.get("table") and s.locator.get("row") == 2]
    anchors = [doc.anchor_for_span(s) for s in original]
    assert [s.text for s in original] == ["84.1", "1.2", "3"]
    assert all(a.table.value_context is not None for a in anchors)
    changed = DocxDocument(
        "p.docx", word_table([("Baseline", "9.2 ± 3.1 (n = 7)"), ("Ours", "80.9 ± 2.4 (n = 5)")])
    )
    assert [changed.locate(a).text for a in anchors] == ["80.9", "2.4", "5"]
    assert all(changed.locate(a).locator["row"] == 3 for a in anchors)
    renamed = DocxDocument("p.docx", word_table([("Ours", "84.1 ± 1.2 (seeds = 3)")]))
    for anchor in anchors:
        with pytest.raises(PaperDeltaError):
            renamed.locate(anchor)


def test_repeated_indistinguishable_cell_contexts_are_not_selected_by_ordinal():
    doc = DocxDocument("p.docx", word_table([("Ours", "Values: 84.1 and 84.1 and 84.1 and 84.1")]))
    spans = [s for s in doc.numbers() if s.text == "84.1"]
    for span in spans[1:3]:
        with pytest.raises(PaperDeltaError):
            doc.anchor_for_span(span)


def test_reviewed_numeric_row_identity_must_resolve_to_chosen_original():
    doc = DocxDocument("p.docx", word_table([("001", "84.1 ± 1.2"), ("002", "84.1 ± 1.2")]))
    target = next(s for s in doc.numbers() if s.text == "84.1")
    with pytest.raises(PaperDeltaError):
        doc.anchor_for_span(target)
    anchor = reviewed_table_anchor(
        doc, target, ReviewedTableIdentity(header_rows=1, row_prefix=["001"])
    )
    assert doc.locate(anchor) == target
    with pytest.raises(PaperDeltaError):
        reviewed_table_anchor(doc, target, ReviewedTableIdentity(header_rows=1, row_prefix=["002"]))


def test_tight_pdf_word_gaps_separate_measurements_without_splitting_identifiers():
    stream = io.BytesIO()
    canvas = Canvas(stream, pagesize=(600, 800), invariant=1)
    canvas.setFont("Helvetica", 10)
    x = 60
    for value in ("Variance", "was", "0.5717", "and", "0.024."):
        canvas.drawString(x, 600, value)
        x += stringWidth(value, "Helvetica", 10) + 1.5
    canvas.drawString(60, 560, "ResNet50 model, run12 identifier")
    canvas.save()
    doc = PdfDocument("p.pdf", stream.getvalue())
    assert [s.text for s in doc.numbers()] == ["0.5717", "0.024"]
    for span in doc.numbers():
        assert doc.locate(doc.anchor_for_span(span)) == span


def test_pdf_two_columns_with_narrow_gutter_keep_separate_contexts():
    stream = io.BytesIO()
    canvas = Canvas(stream, pagesize=(600, 800), invariant=1)
    canvas.setFont("Helvetica", 12)
    left = "Ours 84.1"
    canvas.drawString(50, 600, left)
    right_x = 50 + stringWidth(left, "Helvetica", 12) + 24
    canvas.drawString(right_x, 600, "Baseline 80.9")
    canvas.save()
    doc = PdfDocument("p.pdf", stream.getvalue())
    left, right = doc.numbers()
    assert left.locator["block"] != right.locator["block"]
    assert "Baseline" not in left.context and "Ours" not in right.context


def test_superscript_is_quarantined_without_hiding_unrelated_result():
    stream = io.BytesIO()
    canvas = Canvas(stream, pagesize=(600, 800), invariant=1)
    text = canvas.beginText(40, 720)
    text.setFont("Helvetica", 12)
    text.textOut("Magnitude 10")
    text.setRise(6)
    text.setFont("Helvetica", 8)
    text.textOut("2")
    text.setRise(0)
    text.setFont("Helvetica", 12)
    text.textOut("; accuracy 84.1 percent.")
    canvas.drawText(text)
    canvas.save()
    doc = PdfDocument("p.pdf", stream.getvalue())
    assert [s.text for s in doc.numbers()] == ["84.1"]
    target = doc.numbers()[0]
    assert doc.locate(doc.anchor_for_span(target)) == target
    assert any(i["code"] == "PDF_UNRELIABLE_TEXT" for i in doc.issues)


def test_numeric_context_does_not_expand_to_a_later_repeated_conjunction():
    def document(values):
        stream = io.BytesIO()
        canvas = Canvas(stream, pagesize=(800, 600), invariant=1)
        canvas.drawString(
            30, 500, f"Variance was {values[0]} and {values[1]}. Shares {values[2]}% and 69%"
        )
        canvas.save()
        return PdfDocument("p.pdf", stream.getvalue())

    original = document(("0.5717", "0.024", "31"))
    target = original.numbers()[0]
    anchor = original.anchor_for_span(target)
    assert anchor.numeric_only
    assert original.locate(anchor) == target
    assert document(("0.6828", "0.035", "42")).locate(anchor).text == "0.6828"
    repeated = DocxDocument("p.docx", word_table([("Ours", "84.1 and 84.1 and 84.1 and 84.1")]))
    for span in [s for s in repeated.numbers() if s.text == "84.1"][1:3]:
        with pytest.raises(PaperDeltaError):
            repeated.anchor_for_span(span)


def test_complex_graphics_page_does_not_discard_other_pages():
    stream = io.BytesIO()
    canvas = Canvas(stream, pagesize=(600, 800), invariant=1)
    canvas.drawString(50, 600, "Before 84.1")
    canvas.showPage()
    canvas.drawString(50, 600, "Unverified 99.9")
    for i in range(10001):
        canvas.line(10, 10 + i / 100, 400, 10 + i / 100)
    canvas.showPage()
    canvas.drawString(50, 600, "After 80.9")
    canvas.save()
    doc = PdfDocument("p.pdf", stream.getvalue())
    assert [s.text for s in doc.numbers()] == ["84.1", "80.9"]
    assert doc.unreliable_pages == {2}
    assert any(i["code"] == "PDF_PAGE_LIMIT" for i in doc.issues)
