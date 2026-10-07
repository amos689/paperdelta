"""Whole exponents and static navigation, with partial/active content refusals."""

import io

import pytest
from docx import Document
from pypdf import PdfReader, PdfWriter
from pypdf.generic import ArrayObject, DictionaryObject, NameObject, NumberObject, TextStringObject
from reportlab.pdfgen.canvas import Canvas

from paperdelta.analysis import check_project
from paperdelta.config import config_text
from paperdelta.docx_document import DocxDocument
from paperdelta.errors import PaperDeltaError
from paperdelta.models import Config
from paperdelta.pdf_document import PdfDocument
from paperdelta.storage import Project


def word_scientific(*, exponent="-4", unsafe=None):
    document = Document()
    paragraph = document.add_paragraph("Result 1.2×10")
    run = paragraph.add_run(exponent)
    run.font.superscript = True
    if unsafe == "partial":
        run.text = "-"
        paragraph.add_run("4")
    if unsafe == "hidden":
        run.font.hidden = True
    if unsafe == "whole":
        paragraph.runs[0].font.superscript = True
    paragraph.add_run(" for our method.")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def pdf_scientific(*, unsafe=None):
    output = io.BytesIO()
    canvas = Canvas(output, pagesize=(600, 800), invariant=1)
    text = canvas.beginText(40, 700)
    text.setFont("Helvetica", 12)
    text.textOut("Result 1.2 × 10")
    text.setRise(5 if unsafe != "baseline" else 0)
    text.setFont("Helvetica", 8 if unsafe != "baseline" else 12)
    if unsafe == "hidden":
        text.setTextRenderMode(3)
    text.textOut("-4")
    text.setTextRenderMode(0)
    text.setRise(0)
    text.setFont("Helvetica", 12)
    text.textOut(" for our method.")
    canvas.drawText(text)
    canvas.save()
    return output.getvalue()


@pytest.mark.parametrize("kind", ["docx", "pdf"])
def test_complete_native_scientific_binding_passes_and_then_mismatches(tmp_path, kind):
    raw = word_scientific() if kind == "docx" else pdf_scientific()
    document = (DocxDocument if kind == "docx" else PdfDocument)("p." + kind, raw)
    assert len(document.numbers()) == 1
    span = document.numbers()[0]
    assert document.comparison_text(span) == "1.2e-4"
    with pytest.raises(PaperDeltaError):
        document.validate_numeric_span(document.span(span.start, span.start + 3))
    anchor = document.anchor_for_span(span)
    assert document.locate(anchor) == span
    project = Project(tmp_path)
    project.write("p." + kind, raw)
    project.write("data.csv", b"id,value\none,0.00012\n")
    config = Config.model_validate(
        {
            "schema_version": 10,
            "paper": {"entry": "p." + kind},
            "sources": {
                "data": {
                    "path": "data.csv",
                    "format": "csv",
                    "primary_key": ["id"],
                    "columns": {"id": "string", "value": "decimal"},
                }
            },
            "metrics": {
                "result": {
                    "source": "data",
                    "field": "value",
                    "where": {"id": "one"},
                    "unit": "scalar",
                }
            },
            "occurrences": {
                "result": {
                    "file": "p." + kind,
                    "metric": "result",
                    "anchor": anchor.model_dump(),
                    "display": {"kind": "scientific", "places": 1, "percent_symbol": False},
                }
            },
        }
    )
    project.write("paperdelta.yaml", config_text(config).encode())
    before = check_project(tmp_path)["occurrences"]["result"]
    assert before["status"] == "pass" and before["actual"] == span.text
    project.write("data.csv", b"id,value\none,0.00013\n")
    assert check_project(tmp_path)["occurrences"]["result"]["status"] == "mismatch"
    assert project.read("p." + kind) == raw
    from paperdelta.annotations import preview_annotations

    plan = preview_annotations(project, {"occurrences": ["result"]})
    assert "1.2e-4" in plan["entries"][0]["note"]
    assert plan["entries"][0]["actual"] == span.text


@pytest.mark.parametrize("unsafe", ["partial", "hidden", "whole"])
def test_ambiguous_word_exponents_are_not_flattened_into_plain_numbers(unsafe):
    document = DocxDocument("p.docx", word_scientific(unsafe=unsafe))
    assert not document.numbers() and document.issues


@pytest.mark.parametrize("unsafe", ["baseline", "hidden"])
def test_ambiguous_pdf_exponents_are_not_advertised_as_full_values(unsafe):
    document = PdfDocument("p.pdf", pdf_scientific(unsafe=unsafe))
    assert not document.scientific_ranges
    assert not document.numbers()


@pytest.mark.parametrize(
    "mode", ["array", "goto", "javascript", "next", "remote", "wrong-page", "form"]
)
def test_pdf_static_open_view_is_distinct_from_active_or_invalid_actions(mode):
    output = io.BytesIO()
    canvas = Canvas(output, pagesize=(600, 800), invariant=1)
    canvas.drawString(50, 700, "Result 84.1")
    canvas.save()
    writer = PdfWriter(clone_from=PdfReader(output))
    destination = ArrayObject([writer.pages[0].indirect_reference, NameObject("/Fit")])
    action = DictionaryObject(
        {NameObject("/S"): NameObject("/GoTo"), NameObject("/D"): destination}
    )
    if mode == "javascript":
        action = DictionaryObject(
            {
                NameObject("/S"): NameObject("/JavaScript"),
                NameObject("/JS"): TextStringObject("throw 'not executed'"),
            }
        )
    elif mode == "next":
        action[NameObject("/Next")] = DictionaryObject(
            {NameObject("/S"): NameObject("/JavaScript")}
        )
    elif mode == "remote":
        action[NameObject("/S")] = NameObject("/GoToR")
    elif mode == "wrong-page":
        destination[0] = NumberObject(0)
    elif mode == "form":
        writer.root_object[NameObject("/AcroForm")] = DictionaryObject()
    writer.root_object[NameObject("/OpenAction")] = destination if mode == "array" else action
    copied = io.BytesIO()
    writer.write(copied)
    document = PdfDocument("p.pdf", copied.getvalue())
    assert [s.text for s in document.numbers()] == (["84.1"] if mode in {"array", "goto"} else [])
    assert document.dynamic == (mode not in {"array", "goto"})
