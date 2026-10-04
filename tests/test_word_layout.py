"""Native Word identities survive edits without guessing across merged cells or notes."""

import io
import zipfile
from copy import deepcopy
from xml.etree import ElementTree as ET

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from paperdelta.analysis import check_project
from paperdelta.config import config_text
from paperdelta.docx_document import DocxDocument
from paperdelta.errors import PaperDeltaError
from paperdelta.i18n import language_context, translated
from paperdelta.locations import location_label
from paperdelta.models import Config, TableCellAnchor
from paperdelta.onboarding import init_project
from paperdelta.records import StoredReport, validate_record
from paperdelta.storage import Project

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R = "http://schemas.openxmlformats.org/package/2006/relationships"


def packed(document):
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def notes_document(kind="footnote", *, linked=True, referenced=True, duplicate=False):
    document = Document()
    paragraph = document.add_paragraph("Our accuracy is 84.1%.")
    if referenced:
        run = paragraph.add_run()
        run.font.superscript = True
        ref = OxmlElement("w:" + kind + "Reference")
        ref.set(qn("w:id"), "7")
        run._r.append(ref)
    document.add_paragraph("The baseline is 80.9%.")
    part = (
        f'<w:{kind}s xmlns:w="{W}"><w:{kind} w:id="7">'
        f'<w:p><w:r><w:rPr><w:vertAlign w:val="superscript"/></w:rPr><w:{kind}Ref/>'
        "</w:r><w:r><w:t> Standard error is 0.4.</w:t></w:r></w:p>"
        f'</w:{kind}><w:{kind} w:id="8"><w:p><w:r><w:t>Unused result 99.9.</w:t>'
        f"</w:r></w:p></w:{kind}></w:{kind}s>"
    )
    original = zipfile.ZipFile(io.BytesIO(packed(document)))
    output = io.BytesIO()
    with original, zipfile.ZipFile(output, "w") as archive:
        for name in original.namelist():
            raw = original.read(name)
            if name == "word/_rels/document.xml.rels" and linked:
                root = ET.fromstring(raw)
                for i in range(2 if duplicate else 1):
                    ET.SubElement(
                        root,
                        "{" + R + "}Relationship",
                        {
                            "Id": "notes" + str(i),
                            "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
                            + kind
                            + "s",
                            "Target": kind + "s.xml",
                        },
                    )
                raw = ET.tostring(root)
            archive.writestr(name, raw)
        archive.writestr(f"word/{kind}s.xml", part)
    return output.getvalue()


def result_table(document, caption="Table 1. Accuracy"):
    document.add_paragraph(caption)
    table = document.add_table(rows=4, cols=4)
    table.cell(0, 0).merge(table.cell(0, 3)).text = "Evaluation results"
    for cell, text in zip(
        table.rows[1].cells, ("Dataset", "Model", "95% CI lower", "Upper"), strict=True
    ):
        cell.text = text
    for row, values in zip(
        table.rows[2:],
        (("Test", "Ours", "82.4", "84.2"), ("Test", "Baseline", "80.3", "81.8")),
        strict=True,
    ):
        for cell, text in zip(row.cells, values, strict=True):
            cell.text = text
    table.cell(2, 0).merge(table.cell(3, 0)).text = "Test"
    return table


def test_merged_headers_vertical_labels_and_duplicate_tables_use_literal_identity():
    source = Document()
    first = result_table(source)
    second = result_table(source, "Table 2. Ablation")
    second.cell(2, 2).text = "72.4"
    document = DocxDocument("p.docx", packed(source))
    span = next(s for s in document.numbers() if s.text == "82.4")
    anchor = document.anchor_for_span(span)
    assert anchor.table.header_rows == 2
    assert anchor.table.caption == "Table 1. Accuracy"
    assert anchor.table.row_prefix == ["Test", "Ours"]
    assert span.locator["row"] == 3 and span.locator["cell"] == 3
    assert document.locate(anchor) == span
    first.cell(2, 2).text = "81.4"
    changed = DocxDocument("p.docx", packed(source))
    assert changed.locate(anchor).text == "81.4"
    first.cell(1, 2).text = "Standard deviation"
    with pytest.raises(PaperDeltaError, match="0"):
        DocxDocument("p.docx", packed(source)).locate(anchor)


def test_vertical_owner_is_not_duplicated_and_horizontal_values_are_unknown():
    source = Document()
    table = source.add_table(rows=3, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Group", "Accuracy"
    table.cell(1, 0).text, table.cell(2, 0).text = "Ours", "Ours"
    table.cell(1, 1).merge(table.cell(2, 1)).text = "84.1"
    doc = DocxDocument("p.docx", packed(source))
    span = doc.numbers()[0]
    assert len(doc.numbers()) == 1
    assert doc.locate(doc.anchor_for_span(span)).locator["row"] == 2
    source = Document()
    table = source.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "84.1"
    doc = DocxDocument("p.docx", packed(source))
    assert not doc.numbers()
    assert doc.issues


@pytest.mark.parametrize("mutation", ["orphan", "content", "mismatched-width"])
def test_invalid_vertical_merge_never_yields_a_measured_value(mutation):
    source = Document()
    table = result_table(source)
    if mutation == "orphan":
        table.cell(2, 0)._tc.get_or_add_tcPr().find(qn("w:vMerge")).set(qn("w:val"), "continue")
    elif mutation == "content":
        continuation = table._tbl.tr_lst[3].tc_lst[0]
        continuation.p_lst[0].append(OxmlElement("w:r"))
        text = OxmlElement("w:t")
        text.text = "99.9"
        continuation.p_lst[0][-1].append(text)
    else:
        span = OxmlElement("w:gridSpan")
        span.set(qn("w:val"), "2")
        table._tbl.tr_lst[3].tc_lst[0].get_or_add_tcPr().append(span)
    parsed = DocxDocument("p.docx", packed(source))
    assert not [n for n in parsed.numbers() if n.text == "82.4"]
    assert any(i["code"] == "DOCX_COMPLEX_TABLE" for i in parsed.issues)


@pytest.mark.parametrize("kind", ["footnote", "endnote"])
def test_referenced_notes_keep_part_id_and_read_only_report(tmp_path, kind):
    raw = notes_document(kind)
    document = DocxDocument("p.docx", raw)
    assert not document.issues
    assert [s.text for s in document.numbers()] == ["84.1", "80.9", "0.4"]
    note = document.numbers()[2]
    assert note.locator["part"] == f"word/{kind}s.xml"
    assert note.locator["note_id"] == 7 and note.locator["paragraph"] == 1
    assert document.locate(document.anchor_for_span(note)) == note
    with language_context("en"):
        assert f"{kind} ID 7" in location_label(note.to_dict())
    with language_context("zh-CN"):
        assert ("脚注" if kind == "footnote" else "尾注") + " ID 7" in translated(
            location_label(note.to_dict())
        )
    project = Project(tmp_path)
    project.write("p.docx", raw)
    init_project(project, "p.docx", [])
    report = check_project(tmp_path)
    assert report["report_schema_version"] == 7
    validate_record(StoredReport, report, "REPORT_SCHEMA")
    old = deepcopy(report)
    old["report_schema_version"] = 3
    with pytest.raises(PaperDeltaError):
        validate_record(StoredReport, old, "REPORT_SCHEMA")
    assert project.read("p.docx") == raw


@pytest.mark.parametrize("linked,referenced", [(False, True), (True, False), (False, False)])
def test_unreferenced_or_unlinked_notes_never_become_visible_candidates(linked, referenced):
    document = DocxDocument("p.docx", notes_document(linked=linked, referenced=referenced))
    assert "0.4" not in [s.text for s in document.numbers()]
    assert "99.9" not in [s.text for s in document.numbers()]
    if referenced:
        assert "84.1" not in [s.text for s in document.numbers()]
        assert any(i["code"] == "DOCX_UNSUPPORTED_PARAGRAPH" for i in document.issues)


def test_duplicate_note_links_are_rejected_and_exponents_remain_unknown():
    with pytest.raises(PaperDeltaError) as error:
        DocxDocument("p.docx", notes_document(duplicate=True))
    assert error.value.code == "DOCX_PARSE"
    source = Document()
    source.add_paragraph("10").add_run("2").font.superscript = True
    assert not DocxDocument("p.docx", packed(source)).numbers()


@pytest.mark.parametrize("mutation", ["hidden-reference", "unbalanced-field"])
def test_notes_referenced_only_by_unknown_content_are_not_verified(mutation):
    original = zipfile.ZipFile(io.BytesIO(notes_document()))
    stream = io.BytesIO()
    with original, zipfile.ZipFile(stream, "w") as output:
        for name in original.namelist():
            raw = original.read(name)
            if name == "word/document.xml":
                if mutation == "hidden-reference":
                    raw = raw.replace(b'<w:vertAlign w:val="superscript"/>', b"<w:vanish/>")
                else:
                    raw = raw.replace(
                        b"</w:body>",
                        b'<w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r></w:p></w:body>',
                    )
            output.writestr(name, raw)
    doc = DocxDocument("p.docx", stream.getvalue())
    assert "0.4" not in [span.text for span in doc.numbers()]
    assert doc.issues


def test_layout_anchor_schema_is_explicit_and_legacy_serialization_unchanged(tmp_path):
    legacy = TableCellAnchor(headers=["Model & Accuracy"], row_prefix=["Ours"], column=1)
    assert legacy.model_dump() == {
        "headers": ["Model & Accuracy"],
        "row_prefix": ["Ours"],
        "column": 1,
        "percent_symbol": False,
    }
    with pytest.raises(ValueError):
        TableCellAnchor(**legacy.model_dump(), header_rows=2)
    source = Document()
    result_table(source)
    doc = DocxDocument("p.docx", packed(source))
    anchor = doc.anchor_for_span(next(s for s in doc.numbers() if s.text == "82.4"))
    value = {
        "schema_version": 3,
        "paper": {"entry": "p.docx"},
        "coverage_exclusions": {
            "shown": {
                "file": "p.docx",
                "anchor": anchor.model_dump(),
                "reason": "A reported bound",
                "context_hash": "sha256:" + "0" * 64,
            }
        },
    }
    with pytest.raises(ValueError):
        Config.model_validate(value)
    value["schema_version"] = 7
    assert "header_rows: 2" in config_text(Config.model_validate(value))
