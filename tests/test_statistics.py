"""Independent numerical references, missingness and accepted source identities."""

from decimal import Decimal, localcontext

import pytest
from pydantic import ValidationError

from paperdelta import builder
from paperdelta.analysis import check_project
from paperdelta.config import config_text, load_config
from paperdelta.errors import PaperDeltaError
from paperdelta.models import Display, SourceMetric, StatisticalContract
from paperdelta.onboarding import accept_bindings, init_project, scan_project
from paperdelta.records import StoredReport
from paperdelta.statistics import summarize, t_critical
from paperdelta.storage import Project

CONTRACT = {
    "ddof": 1,
    "unit_of_analysis": "independent experimental seed",
    "confidence_interval": {
        "method": "student_t",
        "level": "0.95",
        "assumption": "independent_normal_observations",
    },
}


@pytest.mark.parametrize(
    "df,expected", [(1, "12.706"), (2, "4.303"), (4, "2.776"), (10, "2.228"), (30, "2.042")]
)
def test_nist_student_t_critical_reference_table(df, expected):
    # https://www.itl.nist.gov/div898/handbook/eda/section3/eda3672.htm
    assert abs(Decimal(t_critical(df, "0.95")) - Decimal(expected)) < Decimal("0.0005")


@pytest.mark.parametrize("level", ["0.5", "0.80", "0.95", "0.999"])
def test_df_one_matches_independent_closed_form_cauchy_quantile(level):
    from mpmath import mp

    context = mp.clone()
    context.dps = 100
    reference = context.tan(context.pi * context.mpf(level) / 2)
    assert abs(context.mpf(t_critical(1, level)) - reference) < context.mpf("1e-75")


@pytest.mark.parametrize("level", ["0.5", "0.8", "0.95", "0.999"])
def test_df_two_matches_independent_closed_form(level):
    # For df = 2, F(t) = 1/2 + t/(2*sqrt(t*t + 2)).
    from mpmath import mp

    context = mp.clone()
    context.dps = 100
    coefficient = context.mpf(level)
    expected = coefficient * context.sqrt(2 / (1 - coefficient**2))
    assert abs(context.mpf(t_critical(2, level)) - expected) < context.mpf("1e-75")


@pytest.mark.parametrize("value", [True, False, 1.0, "1"])
def test_ddof_is_a_strict_integer(value):
    with pytest.raises(ValidationError):
        StatisticalContract(ddof=value, unit_of_analysis="seed")


def test_known_sample_population_sd_and_standard_error():
    values = [Decimal(value) for value in [2, 4, 4, 4, 5, 5, 7, 9]]
    population = summarize(values, StatisticalContract(ddof=0, unit_of_analysis="observation"))
    sample = summarize(values, StatisticalContract(ddof=1, unit_of_analysis="observation"))
    assert population["mean"] == "5" and population["sd"] == "2" and population["n"] == 8
    with localcontext() as context:
        context.prec = 70
        assert abs(Decimal(sample["sd"]) ** 2 - Decimal(32) / 7) < Decimal("1e-65")
        assert abs(Decimal(sample["se"]) ** 2 - Decimal(4) / 7) < Decimal("1e-65")
    assert "confidence_interval" not in sample


def test_large_offset_small_variance_and_order_invariance():
    values = [Decimal("1000000000000000000000000." + suffix) for suffix in ["001", "002", "003"]]
    contract = StatisticalContract(ddof=1, unit_of_analysis="measurement")
    result = summarize(values, contract)
    assert Decimal(result["mean"]) == Decimal("1000000000000000000000000.002")
    assert Decimal(result["sd"]) == Decimal("0.001")
    assert result == summarize(list(reversed(values)), contract)


def test_zero_variance_interval_and_explicit_single_observation_convention():
    result = summarize([Decimal("0.8")] * 5, StatisticalContract(**CONTRACT))
    assert result["confidence_interval"]["lower"] == result["confidence_interval"]["upper"] == "0.8"
    assert (
        summarize([Decimal(3)], StatisticalContract(ddof=0, unit_of_analysis="seed"))["sd"] == "0"
    )
    with pytest.raises(PaperDeltaError, match="n must exceed"):
        summarize([Decimal(3)], StatisticalContract(ddof=1, unit_of_analysis="seed"))


@pytest.mark.parametrize(
    "change",
    [
        {"ddof": 0},
        {"unit_of_analysis": " "},
        {
            "confidence_interval": {
                "method": "normal",
                "level": "0.95",
                "assumption": "independent_normal_observations",
            }
        },
        {
            "confidence_interval": {
                "method": "student_t",
                "level": "0.9999",
                "assumption": "independent_normal_observations",
            }
        },
    ],
)
def test_refuse_undeclared_or_incompatible_statistical_conventions(change):
    with pytest.raises(ValidationError):
        StatisticalContract.model_validate({**CONTRACT, **change})


def make_project(tmp_path, text=r"Accuracy: 82.00 \pm 1.58 (n = 5).", values=None):
    project = Project(tmp_path)
    project.write("paper.tex", (text + "\n").encode())
    values = values or ["80", "81", "82", "83", "84"]
    project.write(
        "results.tsv",
        (
            "model\tseed\taccuracy\n"
            + "".join(f"method\t{i}\t{value}\n" for i, value in enumerate(values, 1))
        ).encode(),
    )
    init_project(project, "paper.tex", ["results.tsv"])
    draft = builder.add_source(
        project,
        builder.start_draft(project),
        name="runs",
        path="results.tsv",
        format="tsv",
        primary_key=["model", "seed"],
        columns={"model": "string", "seed": "integer", "accuracy": "decimal"},
    )
    return project, draft


def metric(project, draft):
    return builder.add_metric(
        project,
        draft,
        name="accuracy",
        source="runs",
        field="accuracy",
        unit="percent",
        reduce="statistics",
        where={"model": "method"},
        expected_count=5,
        expected_seeds=["1", "2", "3", "4", "5"],
        statistics=CONTRACT,
    )


def bind(project, draft, component="mean_sd", show_n=True):
    candidates = scan_project(project)["candidates"]
    return builder.add_occurrences(
        project,
        draft,
        metric="accuracy",
        candidate_ids=[item["candidate_id"] for item in candidates],
        names=["abstract"],
        display_kind="decimal",
        places=2,
        percent_symbol=False,
        statistics={"component": component, "show_n": show_n},
        rationale="Explicit mean and sample SD of the five declared seeds, with n.",
    )


def test_accept_compound_and_recheck_changed_seeds_without_editing_manuscript(tmp_path):
    project, draft = make_project(tmp_path)
    draft = bind(project, metric(project, draft))
    proposal = builder.finalize_draft(project, draft)
    accept_bindings(project, proposal, ["occurrences:abstract"])
    before = check_project(project.root)
    assert before["report_schema_version"] == 9
    assert before["occurrences"]["abstract"]["status"] == "pass"
    assert before["coverage"]["unbound_numbers"] == []
    assert before["metrics"]["accuracy"]["statistics"]["n"] == 5
    assert load_config(project)[0].schema_version == 6
    StoredReport.model_validate(before)
    project.write("results.tsv", project.read("results.tsv").replace(b"\t84\n", b"\t79\n"))
    after = check_project(project.root)
    assert after["occurrences"]["abstract"]["status"] == "mismatch"
    assert (
        after["metrics"]["accuracy"]["fingerprint"] != before["metrics"]["accuracy"]["fingerprint"]
    )
    assert project.read("paper.tex") == b"Accuracy: 82.00 \\pm 1.58 (n = 5).\n"


@pytest.mark.parametrize(
    "values,code",
    [(["80", "81", "82", "83"], "SEED_SET"), (["80", "81", "", "83", "84"], "STATISTICS_MISSING")],
)
def test_incomplete_runs_never_become_successful_statistics(tmp_path, values, code):
    project, draft = make_project(tmp_path, values=values)
    with pytest.raises(PaperDeltaError) as error:
        metric(project, draft)
    assert error.value.code == code
    assert load_config(project)[0].metrics == {}


def test_unknown_after_accepted_seed_disappears(tmp_path):
    project, draft = make_project(tmp_path)
    accept_bindings(
        project,
        builder.finalize_draft(project, bind(project, metric(project, draft))),
        ["occurrences:abstract"],
    )
    project.write("results.tsv", project.read("results.tsv").replace(b"method\t3\t82\n", b""))
    report = check_project(project.root)
    assert report["metrics"]["accuracy"]["status"] == "unknown"
    assert report["metrics"]["accuracy"]["error"] == "SEED_SET"
    assert report["exit_code"] == 2


@pytest.mark.parametrize(
    "component,text,show_n",
    [
        ("mean_se", r"Accuracy: 82.00 \pm 0.71.", False),
        ("ci", "Accuracy: [80.04, 83.96] (n = 5).", True),
        ("mean_ci", "Accuracy: 82.00 [80.04, 83.96].", False),
    ],
)
def test_distinct_explicit_statistical_displays(tmp_path, component, text, show_n):
    project, draft = make_project(tmp_path, text=text)
    draft = bind(project, metric(project, draft), component, show_n)
    accept_bindings(project, builder.finalize_draft(project, draft), ["occurrences:abstract"])
    assert check_project(project.root)["occurrences"]["abstract"]["status"] == "pass"


def test_cannot_relabel_standard_deviation_as_standard_error(tmp_path):
    project, draft = make_project(tmp_path)
    draft = bind(project, metric(project, draft), "mean_se")
    accept_bindings(project, builder.finalize_draft(project, draft), ["occurrences:abstract"])
    assert check_project(project.root)["occurrences"]["abstract"]["status"] == "mismatch"


def test_partial_compound_selection_and_unrelated_numbers_refused(tmp_path):
    project, draft = make_project(tmp_path, text=r"Version 7. Accuracy: 82.00 \pm 1.58 (n = 5).")
    draft = metric(project, draft)
    with pytest.raises(PaperDeltaError) as error:
        bind(project, draft)
    assert error.value.code == "STATISTICAL_SELECTION"


def test_legacy_models_keep_same_serialized_shape():
    model = SourceMetric(source="runs", field="accuracy", unit="fraction")
    assert "statistics" not in model.model_dump()
    assert Display().model_dump() == {"kind": "decimal", "places": 1, "percent_symbol": True}


def test_statistical_schema_cannot_be_downgraded(tmp_path):
    project, draft = make_project(tmp_path)
    _, config = builder.resume_draft(project, metric(project, draft))
    project.write(
        "paperdelta.yaml",
        config_text(config).replace("schema_version: 6", "schema_version: 5").encode(),
    )
    assert check_project(project.root)["exit_code"] == 2


@pytest.mark.parametrize("kind", ["tex", "docx", "pdf"])
@pytest.mark.parametrize("table", [False, True])
def test_native_compound_positions_and_whole_expression_repair(tmp_path, kind, table):
    from io import BytesIO

    from paperdelta.repairs import accept_repairs, propose_repairs, scan_repairs

    project = Project(tmp_path)

    def write(prefix="Accuracy", value="82.00"):
        expression = value + (r" \pm " if kind == "tex" else " ± ") + "1.58 (n = 5)"
        if kind == "tex":
            text = (
                "\\begin{tabular}{lr}\nModel & Accuracy \\\\\nmethod & $"
                + expression
                + "$ \\\\\n\\end{tabular}\n"
                if table
                else prefix + ": $" + expression + "$.\n"
            )
            if prefix != "Accuracy" and not table:
                text = prefix + ": \\(" + expression + "\\).\n"
            raw = text.encode()
        elif kind == "docx":
            from docx import Document

            document = Document()
            if table:
                grid = document.add_table(rows=2, cols=2)
                grid.cell(0, 0).text, grid.cell(0, 1).text = "Model", "Accuracy"
                grid.cell(1, 0).text, grid.cell(1, 1).text = "method", expression
            else:
                document.add_paragraph(prefix + ": " + expression + ".")
            stream = BytesIO()
            document.save(stream)
            raw = stream.getvalue()
        else:
            from reportlab.pdfgen import canvas

            stream = BytesIO()
            page = canvas.Canvas(stream, pagesize=(620, 420))
            if table:
                for y in [340, 310, 280]:
                    page.line(30, y, 580, y)
                for x in [30, 180, 580]:
                    page.line(x, 280, x, 340)
                page.drawString(40, 320, "Model")
                page.drawString(190, 320, "Accuracy")
                page.drawString(40, 290, "method")
                page.drawString(190, 290, expression)
            else:
                page.drawString(30, 350, prefix + ": " + expression + ".")
            page.save()
            raw = stream.getvalue()
        project.write("paper." + kind, raw)

    write()
    project.write(
        "results.tsv",
        b"model\tseed\taccuracy\nmethod\t1\t80\nmethod\t2\t81\nmethod\t3\t82\nmethod\t4\t83\nmethod\t5\t84\n",
    )
    init_project(project, "paper." + kind, ["results.tsv"])
    draft = builder.add_source(
        project,
        builder.start_draft(project),
        name="runs",
        path="results.tsv",
        format="tsv",
        primary_key=["model", "seed"],
        columns={"model": "string", "seed": "integer", "accuracy": "decimal"},
    )
    draft = bind(project, metric(project, draft))
    accept_bindings(project, builder.finalize_draft(project, draft), ["occurrences:abstract"])
    report = check_project(project.root)
    assert report["occurrences"]["abstract"]["status"] == "pass"
    anchor = load_config(project)[0].occurrences["abstract"].anchor
    assert (anchor.table is not None) == table
    assert scan_repairs(project)["broken"] == []
    write(value="81.00")
    mismatch = check_project(project.root)
    assert mismatch["occurrences"]["abstract"]["status"] == "mismatch"
    if kind in {"docx", "pdf"}:
        from paperdelta.patches import create_patch

        before = project.read("paper." + kind)
        with pytest.raises(PaperDeltaError) as error:
            create_patch(project, mismatch, ["abstract"])
        assert error.value.code == "DOCUMENT_READ_ONLY"
        assert project.read("paper." + kind) == before
    if not table:
        write(prefix="Test result")
        repair = scan_repairs(project)
        assert len(repair["broken"]) == 1
        first = next(item for item in repair["candidates"] if item["text"] == "82.00")
        proposal = propose_repairs(
            project,
            [
                {
                    "binding": "occurrences:abstract",
                    "candidate_id": first["candidate_id"],
                    "rationale": "The heading changed; the result and seed identity are unchanged.",
                }
            ],
        )
        accepted = accept_repairs(project, proposal, ["occurrences:abstract"])
        assert accepted["report"]["occurrences"]["abstract"]["status"] == "pass"


def test_batch_statistical_template_and_readonly_agent_proposal(tmp_path):
    from paperdelta.batch import build_proposal, create_catalog, inspect_catalog
    from paperdelta.studio_batch import make_template, template_request

    project, draft = make_project(tmp_path)
    request = {
        "source": "runs",
        "fields": ["accuracy"],
        "group_by": ["model"],
        "unit": "percent",
        "reduce": "statistics",
        "expected_count": 5,
        "expected_seeds": ["1", "2", "3", "4", "5"],
        "statistics": CONTRACT,
        "display": {
            "kind": "decimal",
            "places": 2,
            "percent_symbol": False,
            "statistics": {"component": "mean_sd", "show_n": True},
        },
    }
    catalog = create_catalog(project, draft, request)
    preview = inspect_catalog(project, catalog)
    choice = preview["choices"][0]
    proposal = build_proposal(
        project,
        catalog,
        [
            {
                "choice_id": choice["choice_id"],
                "candidate_ids": [item["candidate_id"] for item in preview["locations"]],
                "rationale": "All declared seeds and the single complete result.",
            }
        ],
    )
    assert len(proposal["additions"]["occurrences"]) == 1
    _, config = builder.resume_draft(project, draft)
    template = make_template("stats", config.sources["runs"], catalog["request"])
    assert template["request"]["statistics"] == CONTRACT
    assert template_request(project, draft, template, "runs")["statistics"] == CONTRACT
    assert load_config(project)[0].metrics == {}


def test_statistics_report_bilingual_language_switch_and_semantics(tmp_path):
    from paperdelta.html_report import html_report
    from paperdelta.i18n import language_context

    project, draft = make_project(tmp_path)
    accept_bindings(
        project,
        builder.finalize_draft(project, bind(project, metric(project, draft))),
        ["occurrences:abstract"],
    )
    report = check_project(project.root)
    for language in ["en", "zh-CN"]:
        with language_context(language):
            html = html_report(report)
        assert "independent experimental seed" in html and "mpmath" not in html
        import json
        import re

        translations = json.loads(re.search(r'id="report-locales">(.*?)</script>', html).group(1))
        assert any(
            "统计显著性" in value["zh-CN"] and "statistical significance" in value["en"]
            for value in translations.values()
        )


def test_compound_patch_rederives_preserves_bytes_and_recovers(tmp_path):
    import copy

    from paperdelta.patches import apply_patch, create_patch, preview_patch, recover_transaction
    from paperdelta.storage import fingerprint

    project, draft = make_project(
        tmp_path, text="\ufeff备注：\r\nAccuracy: 82.00 \\pm 1.58 (n = 5).\r"
    )
    before = project.read("paper.tex")
    accept_bindings(
        project,
        builder.finalize_draft(project, bind(project, metric(project, draft))),
        ["occurrences:abstract"],
    )
    project.write("results.tsv", project.read("results.tsv").replace(b"\t84\n", b"\t79\n"))
    report = check_project(project.root)
    expected = report["occurrences"]["abstract"]["suggestion"]["replacement"]
    report["occurrences"]["abstract"]["suggestion"]["replacement"] = "fabricated result"
    patch = create_patch(project, report, ["abstract"])
    assert patch["changes"][0]["replacement"] == expected
    assert expected in preview_patch(project, patch)
    assert project.read("paper.tex") == before
    tampered = copy.deepcopy(patch)
    tampered["changes"][0]["replacement"] = "99.99"
    tampered["patch_id"] = fingerprint({k: v for k, v in tampered.items() if k != "patch_id"})
    with pytest.raises(PaperDeltaError) as error:
        apply_patch(project, tampered)
    assert error.value.code == "PATCH_NOT_DERIVED"
    applied = apply_patch(project, patch)
    assert applied["report"]["occurrences"]["abstract"]["status"] == "pass"
    assert project.read("paper.tex") == before.replace(
        b"82.00 \\pm 1.58 (n = 5)", expected.encode()
    )
    recover_transaction(project, applied["transaction_id"], write=True)
    assert project.read("paper.tex") == before
    project.write("results.tsv", project.read("results.tsv") + b"\n")
    with pytest.raises(PaperDeltaError) as error:
        apply_patch(project, patch)
    assert error.value.code == "STALE_PATCH"


def test_companion_preview_keeps_statistical_schema_and_contract(tmp_path):
    from paperdelta.manuscripts import change_manuscripts

    project, draft = make_project(tmp_path)
    accept_bindings(
        project,
        builder.finalize_draft(project, bind(project, metric(project, draft))),
        ["occurrences:abstract"],
    )
    project.write("supplement.tex", b"Additional method description.\n")
    original = project.read("paperdelta.yaml")
    preview = change_manuscripts(project, action="add", file="supplement.tex")
    assert project.read("paperdelta.yaml") == original
    assert preview["after"]["schema_version"] == 6
    assert preview["after"]["metrics"]["accuracy"]["statistics"] == CONTRACT
    accepted = change_manuscripts(project, action="add", file="supplement.tex", accept=True)
    assert project.read(accepted["backup"]) == original
    assert check_project(project.root)["occurrences"]["abstract"]["status"] == "pass"


def test_written_confidence_level_is_checked_independently(tmp_path):
    project, draft = make_project(tmp_path, text=r"Confidence: 90\%.")
    draft = metric(project, draft)
    candidate = scan_project(project)["candidates"][0]
    draft = builder.add_occurrences(
        project,
        draft,
        metric="accuracy",
        candidate_ids=[candidate["candidate_id"]],
        names=["confidence"],
        display_kind="percent",
        places=0,
        percent_symbol=True,
        statistics={"component": "confidence_level"},
        rationale="The explicitly declared two-sided confidence coefficient.",
    )
    accept_bindings(project, builder.finalize_draft(project, draft), ["occurrences:confidence"])
    state = check_project(project.root)["occurrences"]["confidence"]
    assert state["status"] == "mismatch"
    assert state["suggestion"]["replacement"] == r"95\%"
