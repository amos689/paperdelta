"""Real local sessions/HTTP: identities, native positions and guarded acceptance."""

import http.client
import io
import json
import threading
from decimal import Decimal

import pytest

from paperdelta.analysis import check_project
from paperdelta.config import load_config
from paperdelta.errors import PaperDeltaError
from paperdelta.i18n import catalog
from paperdelta.storage import Project, parse_json
from paperdelta.studio import StudioSession, browser_value, project_files
from paperdelta.studio_server import MAX_BODY, StudioServer


@pytest.fixture(params=["tex", "docx", "pdf"])
def studio_project(tmp_path, request):
    project = Project(tmp_path / "实验 001")
    path = "论文." + request.param
    if request.param == "tex":
        raw = b"Abstract score: 84.1\\%.\nSecond result: 84.1\\%.\n"
    elif request.param == "docx":
        from docx import Document

        document = Document()
        document.add_paragraph("Abstract score: 84.1%.")
        document.add_paragraph("Second result: 84.1%.")
        stream = io.BytesIO()
        document.save(stream)
        raw = stream.getvalue()
    else:
        from reportlab.pdfgen.canvas import Canvas

        stream = io.BytesIO()
        canvas = Canvas(stream, invariant=True)
        canvas.drawString(60, 740, "Abstract score: 84.1%.")
        canvas.drawString(60, 700, "Second result: 84.1%.")
        canvas.save()
        raw = stream.getvalue()
    project.write(path, raw)
    project.write("results.csv", b"model,seed,score\n001,1,0.840\n001,2,0.842\n002,1,0.123\n")
    return project, path


def call(session, action, payload=None, language="en"):
    return session.execute(
        {
            "action": action,
            "payload": payload or {},
            "revision": session.revision,
            "language": language,
        }
    )


def stage_metric(session, path):
    call(session, "initialize", {"paper": path, "data": ["results.csv"]})
    source = call(session, "source-preview", {"path": "results.csv"})["source"]
    assert source["sample"][0]["model"] == "001"
    call(
        session,
        "source",
        {
            "name": "experiment",
            "path": "results.csv",
            "format": "csv",
            "columns": {"model": "string", "seed": "integer", "score": "decimal"},
            "primary_key": ["model", "seed"],
            "source_hash": source["hash"],
        },
    )
    return call(
        session,
        "metric",
        {
            "name": "ours",
            "source": "experiment",
            "field": "score",
            "unit": "fraction",
            "reduce": "mean",
            "where": {"model": "001"},
            "expected_count": 2,
            "expected_seeds": ["1", "2"],
        },
    )["state"]


def stage_locations(session):
    state = call(session, "state")["state"]
    return call(
        session,
        "locations",
        {
            "metric": "ours",
            "candidate_ids": [item["candidate_id"] for item in state["candidates"]],
            "names": ["abstract", "result"],
            "display_kind": "percent",
            "places": 1,
            "percent_symbol": True,
            "rationale": "Model 001; both declared seeds; mean score.",
        },
    )["state"]


def test_native_session_preview_subset_accept_and_read_only_inputs(studio_project):
    project, path = studio_project
    original = project.read(path), project.read("results.csv")
    session = StudioSession(project)
    assert not call(session, "state")["state"]["initialized"]
    state = stage_metric(session, path)
    assert state["metrics"]["ours"]["result"]["value"] == "0.841"
    config_before = project.read("paperdelta.yaml")
    stage_locations(session)
    assert project.read("paperdelta.yaml") == config_before
    preview = call(session, "preview")["state"]["preview"]
    assert all(item["status"] == "pass" for item in preview["items"])
    assert all(item["location"]["file"] == path for item in preview["items"])
    result = call(
        session,
        "accept",
        {"proposal_id": preview["proposal_id"], "selected": ["occurrences:abstract"]},
    )["state"]
    assert result["receipt"]["bindings"] == ["occurrences:abstract"]
    assert project.read(result["receipt"]["backup"]) == config_before
    config, _ = load_config(project)
    assert list(config.occurrences) == ["abstract"]
    assert config.metrics["ours"].where == {"model": "001"}
    assert check_project(project.root)["coverage"]["pass"] == 1
    assert sum(bool(item["bound"]) for item in result["candidates"]) == 1
    assert (project.read(path), project.read("results.csv")) == original
    if path.endswith("pdf"):
        page = call(session, "page", {"file": path, "page": 1})["preview"]["pages"][0]
        assert page["image"].startswith("data:image/png;base64,")
        assert page["hash"] == session.draft["input_hashes"][path]
        assert len(page["boxes"]) == 2
    for language in ("en", "zh-CN"):
        html = call(session, "report", language=language)["html"]
        assert f'<html lang="{language}"' in html


def test_each_stage_undo_download_restore_is_not_acceptance(studio_project):
    project, path = studio_project
    session = StudioSession(project)
    stage_metric(session, path)
    original = project.read("paperdelta.yaml")
    stage_locations(session)
    exported = call(session, "draft-export")["json"]
    state = call(session, "undo")["state"]
    assert not state["additions"]["occurrences"] and state["metrics"]
    state = call(session, "draft-import", {"draft": parse_json(exported)})["state"]
    assert len(state["additions"]["occurrences"]) == 2
    call(session, "refresh")
    assert not session.history
    assert project.read("paperdelta.yaml") == original
    state = call(session, "draft-import", {"draft": parse_json(exported)})["state"]
    assert state["can_undo"] and state["preview"] is None


@pytest.mark.parametrize("changed", ["paper", "data", "config"])
def test_preview_refuses_any_changed_input_before_acceptance(studio_project, changed):
    project, path = studio_project
    session = StudioSession(project)
    stage_metric(session, path)
    stage_locations(session)
    preview = call(session, "preview")["state"]["preview"]
    target = {"paper": path, "data": "results.csv", "config": "paperdelta.yaml"}[changed]
    project.write(target, project.read(target) + b"\n")
    before = project.read("paperdelta.yaml")
    with pytest.raises(PaperDeltaError) as error:
        call(
            session,
            "accept",
            {"proposal_id": preview["proposal_id"], "selected": ["occurrences:abstract"]},
        )
    assert error.value.code == "STALE_DRAFT"
    assert project.read("paperdelta.yaml") == before
    assert call(session, "state")["state"]["stale"]
    assert not project.path(".paperdelta/config-backups").exists()


def test_other_tab_revision_and_unpreviewed_selection_cannot_write(studio_project):
    project, path = studio_project
    session = StudioSession(project)
    old_revision = session.revision
    stage_metric(session, path)
    before = project.read("paperdelta.yaml")
    with pytest.raises(PaperDeltaError) as error:
        session.execute({"action": "refresh", "revision": old_revision})
    assert error.value.code == "STUDIO_REVISION"
    stage_locations(session)
    with pytest.raises(PaperDeltaError) as error:
        call(session, "accept", {"proposal_id": "guessed", "selected": ["occurrences:abstract"]})
    assert error.value.code == "STUDIO_PREVIEW"
    assert project.read("paperdelta.yaml") == before


def test_identity_errors_and_empty_subset_never_save(studio_project):
    project, path = studio_project
    session = StudioSession(project)
    stage_metric(session, path)
    for where, code in (({"model": "1"}, "EMPTY_SELECTION"), ({"model": 1}, "STUDIO_REQUEST")):
        with pytest.raises(PaperDeltaError) as error:
            call(
                session,
                "metric",
                {
                    "name": "bad",
                    "source": "experiment",
                    "field": "score",
                    "unit": "fraction",
                    "reduce": "unique",
                    "where": where,
                },
            )
        assert error.value.code == code
    stage_locations(session)
    preview = call(session, "preview")["state"]["preview"]
    for selection in ([], ["occurrences:missing"], ["occurrences:abstract"] * 2):
        with pytest.raises(PaperDeltaError):
            call(session, "accept", {"proposal_id": preview["proposal_id"], "selected": selection})
    assert not load_config(project)[0].occurrences


def test_exact_decimal_source_and_derived_roundtrip(tmp_path):
    project = Project(tmp_path)
    project.write("paper.tex", b"Score: 1.0.\n")
    value = "0.12345678901234567890123456789"
    project.write("data.csv", f"setting,score\n{value},1\n".encode())
    session = StudioSession(project)
    call(session, "initialize", {"paper": "paper.tex", "data": ["data.csv"]})
    source = call(session, "source-preview", {"path": "data.csv"})["source"]
    call(
        session,
        "source",
        {
            "name": "exact",
            "path": "data.csv",
            "format": "csv",
            "columns": {"setting": "decimal", "score": "decimal"},
            "primary_key": ["setting"],
            "source_hash": source["hash"],
        },
    )
    result = call(
        session,
        "metric",
        {
            "name": "value",
            "source": "exact",
            "field": "score",
            "unit": "scalar",
            "reduce": "unique",
            "where": {"setting": value},
        },
    )["state"]
    assert result["metrics"]["value"]["definition"]["where"]["setting"] == value
    call(
        session,
        "derived",
        {"name": "zero", "operation": "difference", "left": "value", "right": "value"},
    )
    assert call(session, "state")["state"]["metrics"]["zero"]["result"]["value"] == "0"
    exported = call(session, "draft-export")["json"]
    assert parse_json(exported)["additions"]["metrics"]["value"]["where"]["setting"] == Decimal(
        value
    )
    assert browser_value({"decimal": Decimal(value), "key": 9007199254740993}) == {
        "decimal": value,
        "key": "9007199254740993",
    }


def test_discovery_skips_private_and_generated_directories_and_rejects_escape(tmp_path):
    project = Project(tmp_path / "workspace")
    for path in (
        "paper.tex",
        "data/result.csv",
        ".git/config.json",
        ".venv/a.csv",
        "build/report.json",
    ):
        project.write(path, b"test")
    assert set(project_files(project)["paths"]) == {"paper.tex", "data/result.csv"}
    session = StudioSession(project)
    for path in ("../private.json", "C:/private.csv", "/etc/passwd"):
        with pytest.raises(PaperDeltaError) as error:
            call(session, "source-preview", {"path": path})
        assert error.value.code == "UNSAFE_PATH"


@pytest.fixture
def http_studio(tmp_path):
    server = StudioServer(Project(tmp_path))
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        yield server
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=3)


def request(server, method="POST", path="/api", body=None, headers=None):
    connection = http.client.HTTPConnection("127.0.0.1", server.server_port, timeout=3)
    default = {
        "Authorization": "Bearer " + server.token,
        "Origin": server.origin,
        "Content-Type": "application/json",
    }
    if headers:
        default.update(headers)
    raw = json.dumps(body or {"action": "state"}).encode() if not isinstance(body, bytes) else body
    connection.request(method, path, raw if method == "POST" else None, default)
    response = connection.getresponse()
    result = response.status, dict(response.getheaders()), response.read()
    connection.close()
    return result


def test_real_http_shell_is_static_no_arbitrary_file_routes_or_token_leak(http_studio):
    for path in ("/", "/studio.css", "/studio.js", "/strings.json"):
        status, headers, body = request(http_studio, "GET", path)
        assert status == 200
        assert http_studio.token.encode() not in body
        assert headers["Cache-Control"] == "no-store"
        assert "frame-ancestors 'none'" in headers["Content-Security-Policy"]
    for path in ("/paperdelta.yaml", "/../private", "/%2e%2e/private", "//example.com", "/api"):
        assert request(http_studio, "GET", path)[0] == 404
    status, _, raw = request(http_studio)
    assert status == 200 and not json.loads(raw)["state"]["initialized"]


@pytest.mark.parametrize(
    "headers",
    [
        {"Authorization": "Bearer wrong"},
        {"Authorization": "Bearer é"},
        {"Origin": "http://evil.example"},
        {"Origin": "null"},
        {"Host": "attacker.example"},
        {"Sec-Fetch-Site": "cross-site"},
    ],
)
def test_http_rejects_cross_origin_rebinding_and_wrong_session(http_studio, headers):
    assert request(http_studio, headers=headers)[0] == 403


@pytest.mark.parametrize(
    "body,headers",
    [
        (b"{bad json", {}),
        (b'{"action":"state","action":"refresh"}', {}),
        (b"{}", {"Content-Type": "text/plain"}),
        (b"{}", {"Content-Length": str(MAX_BODY + 1)}),
        (b"{}", {"Transfer-Encoding": "chunked"}),
        (b"{}", {"Content-Length": "-1"}),
        (b"\xff", {}),
        (b'{"action":"state","payload":{"path":"private"}}', {}),
    ],
)
def test_http_bad_bodies_are_bounded_diagnostics(http_studio, body, headers):
    status, _, raw = request(http_studio, body=body, headers=headers)
    assert status == 400 and "error" in json.loads(raw)


def test_http_localized_errors_and_no_cors_preflight(http_studio):
    status, headers, raw = request(http_studio, body={"action": "unknown", "language": "zh-CN"})
    assert status == 400 and json.loads(raw)["message"] == catalog("zh-CN")["studio.unknown_action"]
    assert "Access-Control-Allow-Origin" not in headers
    assert request(http_studio, "OPTIONS")[0] == 405


def test_json_pointer_empty_root_and_array_metrics(tmp_path):
    project = Project(tmp_path)
    project.write("paper.tex", b"Average: 84.1\\%.\n")
    project.write("data.json", b"[0.840,0.842]")
    session = StudioSession(project)
    call(session, "initialize", {"paper": "paper.tex"})
    source = call(session, "source-preview", {"path": "data.json"})["source"]
    assert source["sample"] == [
        {"pointer": "/0", "value": "0.840"},
        {"pointer": "/1", "value": "0.842"},
    ]
    call(
        session,
        "source",
        {
            "name": "json_data",
            "path": "data.json",
            "format": "json",
            "source_hash": source["hash"],
        },
    )
    result = call(
        session,
        "metric",
        {
            "name": "average",
            "source": "json_data",
            "field": "",
            "unit": "fraction",
            "reduce": "mean",
            "expected_count": 2,
        },
    )["state"]
    assert result["metrics"]["average"]["result"]["value"] == "0.841"


def test_unhinted_source_changed_after_inspection_is_not_declared(tmp_path):
    project = Project(tmp_path)
    project.write("paper.tex", b"Result: 1.0.\n")
    project.write("data.json", b"1")
    session = StudioSession(project)
    call(session, "initialize", {"paper": "paper.tex"})
    source = call(session, "source-preview", {"path": "data.json"})["source"]
    project.write("data.json", b"2")
    with pytest.raises(PaperDeltaError) as error:
        call(
            session,
            "source",
            {"name": "data", "path": "data.json", "format": "json", "source_hash": source["hash"]},
        )
    assert error.value.code == "STALE_DRAFT"
    assert not call(session, "state")["state"]["sources"]


def test_http_requires_origin_even_with_valid_bearer(http_studio):
    connection = http.client.HTTPConnection("127.0.0.1", http_studio.server_port, timeout=3)
    connection.request(
        "POST",
        "/api",
        '{"action":"state"}',
        {
            "Authorization": "Bearer " + http_studio.token,
            "Content-Type": "application/json",
        },
    )
    response = connection.getresponse()
    assert response.status == 403
    response.read()
    connection.close()
