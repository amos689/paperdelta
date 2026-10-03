"""Native positions are checked against author-labelled OOXML fixtures, not text dumps."""

import io
import zipfile
from copy import deepcopy

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from paperdelta import builder
from paperdelta.agent import AgentSession
from paperdelta.analysis import check_project
from paperdelta.batch import build_proposal, create_catalog, inspect_catalog
from paperdelta.config import config_text, load_config
from paperdelta.coverage import change_scope
from paperdelta.docx_document import DocxDocument
from paperdelta.errors import PaperDeltaError
from paperdelta.i18n import language_context
from paperdelta.models import Anchor, Config
from paperdelta.onboarding import accept_bindings, init_project, scan_project
from paperdelta.patches import create_patch
from paperdelta.records import StoredReport, validate_record
from paperdelta.repairs import accept_repairs, propose_repairs, scan_repairs
from paperdelta.reports import html_report, text_report
from paperdelta.snapshots import create_snapshot
from paperdelta.storage import Project, json_text, parse_json, sha256
from paperdelta.watch import Watcher

VALUES = ["84.1", "80.9", "3.1", "81.0", "2.5"]
PARAGRAPH_CASES = [
    "plain",
    "split-runs",
    "bold",
    "italic",
    "chinese",
    "unicode-minus",
    "scientific",
    "thousands",
    "percent",
    "headings",
    "abstract",
    "caption",
    "hyperlink",
    "tabs",
    "line-breaks",
    "number-at-start",
    "number-at-end",
    "direct-heading",
]
TABLE_CASES = ["table", "table-split-runs", "table-multi-paragraph", "table-chinese"]


def packed(document):
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def labelled_case(case):
    document = Document()
    expected = []
    if case in TABLE_CASES:
        table = document.add_table(rows=6, cols=2)
        table.cell(0, 0).text = "Model" if case != "table-chinese" else "模型"
        table.cell(0, 1).text = "Accuracy" if case != "table-chinese" else "准确率"
        for i, value in enumerate(VALUES):
            table.cell(i + 1, 0).text = "ABCDE"[i]
            paragraph = table.cell(i + 1, 1).paragraphs[0]
            if case == "table-multi-paragraph":
                paragraph.text = "Selected result"
                paragraph = table.cell(i + 1, 1).add_paragraph()
            for fragment in list(value) if case == "table-split-runs" else [value]:
                paragraph.add_run(fragment)
            expected.append(
                {
                    "text": value,
                    "paragraph": 4 + 2 * i if case != "table-multi-paragraph" else 5 + 3 * i,
                    "table": 1,
                    "row": i + 2,
                    "cell": 2,
                }
            )
        return packed(document), expected
    if case in {"headings", "abstract", "direct-heading"}:
        heading = document.add_heading("Abstract" if case == "abstract" else "Results", 1)
        if case == "direct-heading":
            heading.style = "Normal"
            outline = OxmlElement("w:outlineLvl")
            outline.set(qn("w:val"), "0")
            heading._p.get_or_add_pPr().append(outline)
    for i, value in enumerate(VALUES):
        if case == "unicode-minus":
            value = "−" + value
        elif case == "scientific":
            value += "e-3"
        elif case == "thousands":
            value = ["1,084.1", "1,080.9", "1,003.1", "1,081.0", "1,002.5"][i]
        paragraph = document.add_paragraph(style="Caption" if case == "caption" else None)
        prefix = f"Result {'ABCDE'[i]}: "
        suffix = "%." if case == "percent" else "."
        if case == "chinese":
            prefix, suffix = f"模型{'甲乙丙丁戊'[i]}准确率", "个百分点。"
        if case == "number-at-start":
            prefix, suffix = "", f" is the result for {'ABCDE'[i]}."
        elif case == "number-at-end":
            suffix = ""
        paragraph.add_run(prefix)
        if case in {"tabs", "line-breaks"}:
            run = paragraph.add_run()
            run.add_tab() if case == "tabs" else run.add_break()
        if case == "hyperlink":
            link = OxmlElement("w:hyperlink")
            run, text = OxmlElement("w:r"), OxmlElement("w:t")
            text.text = value
            run.append(text)
            link.append(run)
            paragraph._p.append(link)
        else:
            for fragment in list(value) if case == "split-runs" else [value]:
                run = paragraph.add_run(fragment)
                run.bold = case == "bold"
                run.italic = case == "italic"
        paragraph.add_run(suffix)
        expected.append(
            {
                "text": value,
                "paragraph": i + (2 if case in {"headings", "abstract", "direct-heading"} else 1),
            }
        )
    return packed(document), expected


@pytest.mark.parametrize("case", PARAGRAPH_CASES + TABLE_CASES)
def test_labelled_positions_and_supported_structure(case):
    raw, expected = labelled_case(case)
    document = DocxDocument("manuscript.docx", raw)
    assert not document.issues
    spans = document.numbers()
    assert len(spans) == 5
    for span, label in zip(spans, expected, strict=True):
        assert span.text == label["text"]
        for key, value in label.items():
            if key != "text":
                assert span.locator[key] == value
        value = span.to_dict()
        assert not {"line", "page", "byte_start", "byte_end"} & value.keys()
        assert value["format"] == "docx"
        assert value["parser"].startswith("paperdelta-docx/1;")
        anchor = builder.anchor_for_span(document, span)
        assert document.locate(anchor).to_dict() == value
    if case == "abstract":
        assert len(document.priority_regions) == 6
    if case == "caption":
        assert all(span.locator["style"].lower() == "caption" for span in spans)
    assert document.hash == sha256(raw)


@pytest.mark.parametrize(
    "case",
    [
        "revision",
        "field",
        "equation",
        "textbox",
        "hidden",
        "symbol",
        "content-control",
        "merged",
        "nested",
        "note",
    ],
)
def test_unsupported_structures_never_create_checkable_numbers(case):
    document = Document()
    if case in {"merged", "nested"}:
        table = document.add_table(rows=2, cols=2)
        if case == "merged":
            table.cell(0, 0).merge(table.cell(0, 1)).text = "Result 84.1"
        else:
            table.cell(0, 0).add_table(rows=1, cols=1).cell(0, 0).text = "Result 84.1"
    else:
        paragraph = document.add_paragraph("Result ")
        run = paragraph.add_run("84.1")
        if case == "revision":
            revision = OxmlElement("w:ins")
            paragraph._p.remove(run._r)
            revision.append(run._r)
            paragraph._p.append(revision)
        elif case == "field":
            field = OxmlElement("w:fldSimple")
            field.set(qn("w:instr"), "= 84.1")
            paragraph._p.append(field)
        elif case == "equation":
            paragraph._p.append(OxmlElement("m:oMath"))
        elif case == "textbox":
            run._r.append(OxmlElement("w:txbxContent"))
        elif case == "hidden":
            run.font.hidden = True
        elif case == "symbol":
            run._r.append(OxmlElement("w:sym"))
        elif case == "note":
            run._r.append(OxmlElement("w:footnoteReference"))
        else:
            paragraph._p.append(OxmlElement("w:sdt"))
    parsed = DocxDocument("paper.docx", packed(document))
    assert parsed.issues
    assert not parsed.numbers()
    with pytest.raises(PaperDeltaError):
        parsed.locate(Anchor(exact="84.1"))


def test_numeric_changes_survive_but_identical_paragraphs_and_semantic_changes_require_repair():
    original = Document()
    original.add_paragraph("Our accuracy is 84.1%.")
    document = DocxDocument("paper.docx", packed(original))
    span = builder.candidate_span(
        document, document.numbers()[0].to_dict(), builder.Display(kind="percent")
    )
    anchor = builder.anchor_for_span(document, span)
    original.paragraphs[0].text = "Our accuracy is 80.9%."
    original.add_paragraph("An unrelated introductory paragraph.")
    changed = DocxDocument("paper.docx", packed(original))
    assert changed.locate(anchor).text == "80.9%"
    original.add_paragraph("Our accuracy is 80.9%.")
    duplicate = DocxDocument("paper.docx", packed(original))
    with pytest.raises(PaperDeltaError, match="2") as error:
        duplicate.locate(anchor)
    assert error.value.code == "ANCHOR_AMBIGUOUS"
    original.paragraphs[0].text = "Baseline accuracy is 80.9%."
    original.paragraphs[2].text = "Baseline accuracy is 80.9%."
    with pytest.raises(PaperDeltaError) as error:
        DocxDocument("paper.docx", packed(original)).locate(anchor)
    assert error.value.code == "ANCHOR_MISSING"


def test_read_only_binding_check_snapshot_percent_and_bilingual_report(tmp_path):
    project = Project(tmp_path)
    document = Document()
    document.add_heading("Abstract", 1)
    paragraph = document.add_paragraph("Our accuracy is ")
    paragraph.add_run("84.").bold = True
    paragraph.add_run("1%.")
    original = packed(document)
    project.write("paper.docx", original)
    project.write("results.csv", b"model,score\nOurs,0.809\n")
    init_project(project, "paper.docx", ["results.csv"])
    assert load_config(project)[0].schema_version == 3
    draft = builder.start_draft(project)
    draft = builder.add_source(
        project,
        draft,
        name="results",
        path="results.csv",
        format="csv",
        columns={"model": "string", "score": "decimal"},
        primary_key=["model"],
    )
    draft = builder.add_metric(
        project,
        draft,
        name="accuracy",
        source="results",
        field="score",
        unit="fraction",
        reduce="unique",
        where={"model": "Ours"},
    )
    draft = builder.add_occurrences(
        project,
        draft,
        metric="accuracy",
        candidate_ids=[scan_project(project)["candidates"][0]["candidate_id"]],
        names=["abstract_accuracy"],
        display_kind="percent",
        places=1,
        percent_symbol=True,
        rationale="Ours accuracy in the abstract.",
    )
    proposal = builder.finalize_draft(project, parse_json(json_text(draft)))
    accept_bindings(project, proposal, list(proposal["rationale"]))
    report = check_project(tmp_path)
    assert report["report_schema_version"] == 3
    assert report["exit_code"] == 1
    assert report["occurrences"]["abstract_accuracy"]["expected"] == "80.9%"
    assert any(a["kind"] == "update_document" for a in report["actions"])
    create_snapshot(project, "before", report)
    validate_record(StoredReport, parse_json(json_text(report)), "REPORT_SCHEMA")
    for selection in (None, ["abstract_accuracy"]):
        with pytest.raises(PaperDeltaError) as error:
            create_patch(project, report, selection)
        assert error.value.code == "DOCUMENT_READ_ONLY"
    for language, label in (("en", "paragraph 2"), ("zh-CN", "第 2 段")):
        with language_context(language):
            assert label in text_report(report)
            html = html_report(report)
            assert label in html and "Our accuracy is 84.1%." in html
    assert project.read("paper.docx") == original
    document.paragraphs[1].text = "Our accuracy is 80.9%."
    project.write("paper.docx", packed(document))
    assert check_project(tmp_path)["exit_code"] == 0


def test_side_parts_and_malformed_packages_are_explicit():
    document = Document()
    document.add_paragraph("Result 84.1.")
    document.sections[0].header.paragraphs[0].text = "Header result 80.9."
    assert "DOCX_UNSUPPORTED_PART" in {
        i["code"] for i in DocxDocument("p.docx", packed(document)).issues
    }
    for raw in (b"not a zip", b"PK\x03\x04"):
        with pytest.raises(PaperDeltaError) as error:
            DocxDocument("p.docx", raw)
        assert error.value.code == "DOCX_PARSE"
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        archive.writestr("word/document.xml", '<!DOCTYPE doc [<!ENTITY x "84.1">]><doc>&x;</doc>')
    with pytest.raises(PaperDeltaError) as error:
        DocxDocument("p.docx", stream.getvalue())
    assert error.value.code == "DOCX_PARSE"


def test_native_schema_and_locations_cannot_claim_byte_addresses(tmp_path):
    document = Document()
    document.add_paragraph("Result 84.1.")
    project = Project(tmp_path)
    project.write("paper.docx", packed(document))
    init_project(project, "paper.docx", [])
    report = check_project(tmp_path)
    malformed = deepcopy(report)
    malformed["coverage"]["unbound_numbers"][0]["byte_start"] = 0
    with pytest.raises(PaperDeltaError):
        validate_record(StoredReport, malformed, "REPORT_SCHEMA")
    malformed = deepcopy(report)
    malformed["report_schema_version"] = 2
    with pytest.raises(PaperDeltaError):
        validate_record(StoredReport, malformed, "REPORT_SCHEMA")
    with pytest.raises(ValueError):
        Config(schema_version=2, paper={"entry": "paper.docx"})
    config, _ = load_config(project)
    value = config.model_dump()
    value["paper"]["macros"] = {"accuracy": 1}
    project.write("paperdelta.yaml", config_text(Config.model_validate(value)).encode())
    assert check_project(tmp_path)["diagnostics"][0]["rule"] == "DOCUMENT_MACROS"


@pytest.fixture
def review_project(tmp_path):
    project = Project(tmp_path)
    document = Document()
    document.add_heading("Abstract", 1)
    document.add_paragraph("Our accuracy is 84.1%.")
    document.add_heading("Results", 1)
    document.add_paragraph("We outperform Baseline.")
    document.add_paragraph("Sample size is 20.")
    project.write("paper.docx", packed(document))
    project.write("results.csv", b"model,score\nOurs,0.841\nBaseline,0.810\n")
    native = DocxDocument("paper.docx", project.read("paper.docx"))
    span = native.numbers()[0]
    anchor = builder.anchor_for_span(native, native.span(span.start, span.end + 1))
    config = Config.model_validate(
        {
            "schema_version": 3,
            "paper": {"entry": "paper.docx"},
            "sources": {
                "results": {
                    "path": "results.csv",
                    "format": "csv",
                    "primary_key": ["model"],
                    "columns": {"model": "string", "score": "decimal"},
                }
            },
            "metrics": {
                key: {
                    "source": "results",
                    "field": "score",
                    "where": {"model": label},
                    "unit": "fraction",
                    "expected_count": 1,
                }
                for key, label in (("accuracy", "Ours"), ("baseline", "Baseline"))
            },
            "occurrences": {
                "accuracy": {
                    "file": "paper.docx",
                    "metric": "accuracy",
                    "anchor": anchor.model_dump(),
                    "display": {"kind": "percent"},
                }
            },
            "claims": {
                "comparison": {
                    "file": "paper.docx",
                    "anchor": {"exact": "We outperform Baseline."},
                    "predicate": {"op": "greater_than", "left": "accuracy", "right": "baseline"},
                }
            },
        }
    )
    project.write("paperdelta.yaml", config_text(config).encode())
    return project, document


def test_word_scope_claims_repair_and_stale_draft(review_project):
    project, document = review_project
    before = project.read("paper.docx")
    assert check_project(project.root)["exit_code"] == 0
    change_scope(project, action="set", regions=["abstract"], require_complete=True, accept=True)
    config, _ = load_config(project)
    assert config.schema_version == 3
    report = check_project(project.root)
    assert len(report["coverage"]["outside_scope_numbers"]) == 1
    assert report["exit_code"] == 0
    create_snapshot(project, "before", report)
    draft = builder.start_draft(project)
    project.write("results.csv", project.read("results.csv").replace(b"0.841", b"0.809"))
    report = check_project(project.root)
    assert report["claims"]["comparison"]["status"] == "mismatch"
    assert report["occurrences"]["accuracy"]["suggestion"]["blocked_by"] == ["claim:comparison"]
    with pytest.raises(PaperDeltaError) as error:
        builder.finalize_draft(project, draft)
    assert error.value.code == "STALE_DRAFT"
    assert project.read("paper.docx") == before
    document.paragraphs[1].text = "Our updated accuracy is 80.9%."
    updated = packed(document)
    project.write("paper.docx", updated)
    scan = scan_repairs(project, baseline="before")
    candidate = next(item for item in scan["candidates"] if item["text"] == "80.9")
    proposal = propose_repairs(
        project,
        [
            {
                "binding": "occurrences:accuracy",
                "candidate_id": candidate["candidate_id"],
                "rationale": "Same Ours accuracy; reworded abstract.",
            }
        ],
        baseline="before",
    )
    accept_repairs(project, proposal, ["occurrences:accuracy"])
    assert check_project(project.root)["occurrences"]["accuracy"]["status"] == "pass"
    assert project.read("paper.docx") == updated


def test_word_watch_recovers_partial_save_and_agent_views_are_read_only(review_project):
    project, document = review_project
    watcher = Watcher(project, debounce=0)
    assert watcher.step(0)["exit_code"] == 0
    project.write("paper.docx", b"PK partial Word save")
    assert watcher.step(1)["exit_code"] == 2
    document.paragraphs[1].text = "Our accuracy is 80.9%."
    final = packed(document)
    project.write("paper.docx", final)
    assert watcher.step(2)["coverage"]["mismatch"] == 1
    agent = AgentSession(project)
    for language in ("en", "zh-CN"):
        with language_context(language):
            report = agent.check_project()
            assert report["occurrences"]["accuracy"]["location"]["locator"]["paragraph"] == 2
            assert agent.scan_project()["candidates"][0]["format"] == "docx"
    watcher.stop()
    assert project.read("paper.docx") == final


def test_word_batch_uses_column_and_row_identity_and_survives_row_reorder(tmp_path):
    project = Project(tmp_path)
    raw, _ = labelled_case("table")
    project.write("paper.docx", raw)
    project.write(
        "results.csv",
        (
            "model,Accuracy\n"
            + "\n".join(f"{name},{value}" for name, value in zip("ABCDE", VALUES, strict=True))
        ).encode(),
    )
    init_project(project, "paper.docx", ["results.csv"])
    draft = builder.add_source(
        project,
        builder.start_draft(project),
        name="results",
        path="results.csv",
        format="csv",
        columns={"model": "string", "Accuracy": "decimal"},
        primary_key=["model"],
    )
    catalog = create_catalog(
        project,
        draft,
        {
            "source": "results",
            "fields": ["Accuracy"],
            "group_by": ["model"],
            "unit": "scalar",
            "reduce": "unique",
            "expected_count": 1,
            "display": {"kind": "decimal", "places": 1},
        },
    )
    preview = inspect_catalog(project, catalog)
    selections = []
    for choice in preview["choices"]:
        model = choice["definition"]["where"]["model"]
        location = next(
            item for item in preview["locations"] if item["row"].split("|")[0].strip() == model
        )
        assert location["column_header"] == "Accuracy"
        selections.append(
            {
                "choice_id": choice["choice_id"],
                "candidate_ids": [location["candidate_id"]],
                "rationale": "Explicit model and accuracy column.",
            }
        )
    proposal = build_proposal(project, catalog, selections)
    accept_bindings(project, proposal, list(proposal["rationale"]))
    assert check_project(tmp_path)["coverage"]["pass"] == 5
    document = Document(io.BytesIO(raw))
    table = document.tables[0]
    row = table.rows[1]._tr
    table._tbl.remove(row)
    table._tbl.append(row)
    changed = packed(document)
    project.write("paper.docx", changed)
    assert check_project(tmp_path)["coverage"]["pass"] == 5
    table.cell(0, 1).text = "Precision"
    project.write("paper.docx", packed(document))
    assert check_project(tmp_path)["coverage"]["unknown"] == 5


def test_word_mcp_preserves_native_positions_without_writes(review_project):
    import asyncio

    mcp = pytest.importorskip("mcp")
    from paperdelta.mcp_server import create_server

    project, _ = review_project
    before = {
        p.relative_to(project.root): p.read_bytes() for p in project.root.rglob("*") if p.is_file()
    }

    async def run():
        async with mcp.Client(create_server(project)) as client:
            result = await client.call_tool("check_project")
            assert not result.is_error
            location = result.structured_content["occurrences"]["accuracy"]["location"]
            assert location["format"] == "docx" and "byte_start" not in location
            result = await client.call_tool("scan_project")
            assert not result.is_error
            assert result.structured_content["candidates"][0]["format"] == "docx"

    asyncio.run(run())
    assert before == {
        p.relative_to(project.root): p.read_bytes() for p in project.root.rglob("*") if p.is_file()
    }


@pytest.mark.parametrize("language", ["en", "zh-CN"])
@pytest.mark.parametrize("scenario", ["baseline", "changed", "safe-update"])
def test_bundled_word_demo(tmp_path, language, scenario):
    from paperdelta.demo import create_demo

    with language_context(language):
        result = create_demo(Project(tmp_path), "word-demo", scenario, "docx")
    assert result["check_exit_code"] == (0 if scenario == "baseline" else 1)
    assert result["patch"] is None
    project = Project(tmp_path / "word-demo")
    before = parse_json(project.text("before/report.json")[0])
    assert before["coverage"]["pass"] == 4
    assert (
        check_project(project.root)["input_hashes"]["paper.docx"]
        == before["input_hashes"]["paper.docx"]
    )
    assert f'<html lang="{language}"' in project.text("review/report.html")[0]


def test_optional_parser_is_diagnosed_before_creating_demo(tmp_path, monkeypatch):
    import builtins

    from paperdelta.demo import create_demo

    original = builtins.__import__

    def without_docx(name, *args, **kwargs):
        if name == "docx":
            raise ImportError("optional parser intentionally absent")
        return original(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", without_docx)
    with pytest.raises(PaperDeltaError) as error:
        create_demo(Project(tmp_path), "word-demo", document="docx")
    assert error.value.code == "DOCUMENT_DEPENDENCY"
    assert "paperdelta[docx]" in str(error.value)
    assert not (tmp_path / "word-demo").exists()


@pytest.mark.parametrize("case", ["hidden-style", "numbering", "superscript"])
def test_generated_or_styled_content_is_not_silently_verified(case):
    from docx.enum.style import WD_STYLE_TYPE

    document = Document()
    paragraph = document.add_paragraph("Result 84.1.")
    if case == "hidden-style":
        style = document.styles.add_style("HiddenResult", WD_STYLE_TYPE.PARAGRAPH)
        style.font.hidden = True
        paragraph.style = style
    elif case == "numbering":
        paragraph.style = "List Number"
    else:
        paragraph.runs[0].font.superscript = True
    parsed = DocxDocument("paper.docx", packed(document))
    assert parsed.issues and not parsed.numbers()


def test_inherited_heading_keeps_the_actual_section_and_default_tabs_are_not_text():
    from docx.enum.style import WD_STYLE_TYPE
    from docx.shared import Inches

    document = Document()
    style = document.styles.add_style("ResearchSection", WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = document.styles["Heading 1"]
    document.add_paragraph("Abstract", style=style)
    paragraph = document.add_paragraph("Result 84.1.")
    paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(1))
    parsed = DocxDocument("paper.docx", packed(document))
    assert not parsed.issues
    assert parsed.numbers()[0].locator["section"] == ["Abstract"]
    assert "\t" not in parsed.text
    assert len(parsed.priority_regions) == 2
