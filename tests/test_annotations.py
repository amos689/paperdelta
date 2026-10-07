"""Original bytes, existing comments and rendered positions survive copy export."""

import io
from copy import deepcopy
from zipfile import ZipFile

import pypdfium2 as pdfium
import pytest
from docx import Document
from lxml import etree
from pypdf import PdfReader

from paperdelta.annotations import annotation_bytes, preview_annotations
from paperdelta.cli import main
from paperdelta.demo import create_demo
from paperdelta.docx_document import DocxDocument
from paperdelta.errors import PaperDeltaError
from paperdelta.i18n import language_context
from paperdelta.native_comments import pdf_copy, word_copy
from paperdelta.pdf_document import PdfDocument
from paperdelta.storage import Project, parse_json


def packed(document):
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


@pytest.mark.parametrize("kind", ["docx", "pdf"])
@pytest.mark.parametrize("language", ["en", "zh-CN"])
def test_bilingual_preview_copies_keep_originals_and_claim_attention(tmp_path, kind, language):
    create_demo(Project(tmp_path), "paper", document=kind)
    project = Project(tmp_path / "paper")
    before = {
        p.relative_to(project.root).as_posix(): p.read_bytes()
        for p in project.root.rglob("*")
        if p.is_file()
    }
    with language_context(language):
        plan = preview_annotations(
            project, {"occurrences": ["abstract_accuracy", "table_accuracy"]}
        )
    assert not plan["refused"] and len(plan["entries"]) == 2
    assert len(plan["copies"]) == 1
    assert ("预期显示" if language == "zh-CN" else "Expected display") in plan["entries"][0]["note"]
    assert "claim:main_comparison" in plan["entries"][0]["note"]
    with ZipFile(io.BytesIO(annotation_bytes(project, plan))) as archive:
        manifest = parse_json(archive.read("review.json").decode())
        assert manifest["preview_id"] == plan["preview_id"]
        assert len(manifest["copies_sha256"]) == 1
        copied = archive.read(plan["copies"][0]["path"])
    original = project.read(plan["copies"][0]["original"])
    if kind == "docx":
        old = DocxDocument("paper.docx", original)
        new = DocxDocument("paper.docx", copied)
        assert new.text == old.text
        with ZipFile(io.BytesIO(copied)) as archive:
            comments = etree.fromstring(archive.read("word/paperdelta-comments.xml"))
            assert len(comments) == 2
            assert plan["entries"][0]["note"].splitlines()[0] in "".join(comments.itertext())
    else:
        old, new = PdfReader(io.BytesIO(original)), PdfReader(io.BytesIO(copied))
        assert [p.extract_text() for p in new.pages] == [p.extract_text() for p in old.pages]
        notes = [a.get_object() for page in new.pages for a in page.get("/Annots", [])]
        assert len([a for a in notes if a["/Subtype"] == "/Highlight"]) == 2
    assert before == {
        p.relative_to(project.root).as_posix(): p.read_bytes()
        for p in project.root.rglob("*")
        if p.is_file()
    }


@pytest.mark.parametrize("change", ["evidence", "manuscript", "forged_note", "forged_position"])
def test_export_rejects_changed_inputs_or_forged_preview(tmp_path, change):
    create_demo(Project(tmp_path), "paper", document="docx")
    project = Project(tmp_path / "paper")
    plan = preview_annotations(project, {"occurrences": ["abstract_accuracy"]})
    if change == "evidence":
        project.write(
            "results/metrics.csv", project.read("results/metrics.csv").replace(b"0.807", b"0.700")
        )
    elif change == "manuscript":
        doc = Document(io.BytesIO(project.read("paper.docx")))
        doc.add_paragraph("Later author text.")
        project.write("paper.docx", packed(doc))
    elif change == "forged_note":
        plan["entries"][0]["note"] = "Everything is verified."
    else:
        plan["entries"][0]["location"]["locator"]["offset"] += 1
    with pytest.raises(PaperDeltaError):
        annotation_bytes(project, plan)


def test_word_split_runs_preserve_format_existing_comments_and_parts():
    source = Document()
    paragraph = source.add_paragraph()
    paragraph.add_run("Before 8").bold = True
    paragraph.add_run("4.").italic = True
    paragraph.add_run("1 after").underline = True
    source.add_comment(paragraph.runs[0], text="Existing reviewer note", author="Reviewer")
    raw = packed(source)
    document = DocxDocument("p.docx", raw)
    start = document.text.index("84.1")
    span = document.span(start, start + 4)
    copied = word_copy(
        document, [{"location": span.to_dict(), "actual": span.text, "note": "New review"}]
    )
    with ZipFile(io.BytesIO(raw)) as old, ZipFile(io.BytesIO(copied)) as new:
        for name in old.namelist():
            if name not in {
                "word/document.xml",
                "word/comments.xml",
                "word/_rels/document.xml.rels",
                "[Content_Types].xml",
            }:
                assert old.read(name) == new.read(name), name
        comments = etree.fromstring(new.read("word/comments.xml"))
        assert len(comments) == 2 and "Existing reviewer note" in "".join(comments[0].itertext())
    loaded = Document(io.BytesIO(copied))
    assert loaded.paragraphs[0].text == "Before 84.1 after"
    letters = [
        (c, r.bold, r.italic, r.underline) for r in loaded.paragraphs[0].runs for c in r.text
    ]
    original = [
        (c, r.bold, r.italic, r.underline) for r in source.paragraphs[0].runs for c in r.text
    ]
    assert letters == original
    assert DocxDocument("p.docx", copied).text == document.text


def test_pdf_copy_preserves_existing_reviewer_annotation():
    from pypdf import PdfWriter
    from pypdf.annotations import Text
    from test_pdf_layout import geometry_pdf

    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(geometry_pdf())))
    writer.add_annotation(0, Text(rect=(200, 450, 240, 490), text="Existing review"))
    stream = io.BytesIO()
    writer.write(stream)
    document = PdfDocument("p.pdf", stream.getvalue())
    span = document.numbers()[0]
    copied = pdf_copy(
        document, [{"location": span.to_dict(), "actual": span.text, "note": "New review"}]
    )
    annotations = [a.get_object() for a in PdfReader(io.BytesIO(copied)).pages[0]["/Annots"]]
    assert len(annotations) == 2
    assert annotations[0]["/Contents"] == "Existing review"
    assert list(annotations[0]["/Rect"]) == [200, 450, 240, 490]
    assert annotations[1]["/Contents"] == "New review"


def test_note_comment_refusal_does_not_hide_exportable_body_selection(tmp_path):
    from test_word_layout import notes_document

    from paperdelta.config import config_text
    from paperdelta.models import Config

    project = Project(tmp_path)
    raw = notes_document()
    project.write("p.docx", raw)
    project.write("data.csv", b"id,value\none,84.1\n")
    document = DocxDocument("p.docx", raw)
    body, _, note = document.numbers()
    config = Config.model_validate(
        {
            "schema_version": 10,
            "paper": {"entry": "p.docx"},
            "sources": {
                "data": {
                    "path": "data.csv",
                    "format": "csv",
                    "primary_key": ["id"],
                    "columns": {"id": "string", "value": "decimal"},
                }
            },
            "metrics": {"result": {"source": "data", "field": "value", "unit": "scalar"}},
            "occurrences": {
                name: {
                    "file": "p.docx",
                    "metric": "result",
                    "anchor": document.anchor_for_span(span).model_dump(),
                    "display": {"places": 1},
                }
                for name, span in (("body", body), ("note", note))
            },
        }
    )
    project.write("paperdelta.yaml", config_text(config).encode())
    plan = preview_annotations(project, {"occurrences": ["body", "note"]})
    assert [e["occurrence"] for e in plan["entries"]] == ["body"]
    assert [r["code"] for r in plan["refused"]] == ["ANNOTATION_PART"]
    with ZipFile(io.BytesIO(annotation_bytes(project, plan))) as archive:
        copied = archive.read(plan["copies"][0]["path"])
    with ZipFile(io.BytesIO(raw)) as old, ZipFile(io.BytesIO(copied)) as new:
        assert old.read("word/footnotes.xml") == new.read("word/footnotes.xml")
    assert project.read("p.docx") == raw


@pytest.mark.parametrize("rotation", [0, 90, 180, 270])
@pytest.mark.parametrize("offset", [False, True])
def test_pdf_highlight_raster_matches_original_rotated_cropped_value(rotation, offset):
    from test_pdf_layout import geometry_pdf

    raw = geometry_pdf(rotation, offset)
    document = PdfDocument("p.pdf", raw)
    span = document.numbers()[0]
    copied = pdf_copy(
        document, [{"location": span.to_dict(), "actual": span.text, "note": "Value review"}]
    )
    old, new = PdfReader(io.BytesIO(raw)), PdfReader(io.BytesIO(copied))
    for key in ("/MediaBox", "/CropBox", "/Rotate"):
        assert new.pages[0][key] == old.pages[0][key]
    assert new.pages[0].get_contents().get_data() == old.pages[0].get_contents().get_data()
    images = []
    for content in (raw, copied):
        with pdfium.PdfDocument(content) as pdf:
            images.append(pdf[0].render(scale=2, draw_annots=True).to_pil().convert("RGB"))
    assert images[0].size == images[1].size
    changed = []
    before, after = images[0].load(), images[1].load()
    for y in range(images[0].height):
        for x in range(images[0].width):
            if before[x, y] != after[x, y]:
                changed.append((x / 2, y / 2))
    assert len(changed) > 30
    box = list(map(float, span.locator["bbox"]))
    # PDFium's highlight outline has a small rounded margin around the glyph box.
    assert all(box[0] - 4 <= x <= box[2] + 4 and box[1] - 4 <= y <= box[3] + 4 for x, y in changed)


def test_cli_requires_review_and_never_overwrites_outputs(tmp_path, capsys):
    create_demo(Project(tmp_path), "paper", document="docx")
    root = tmp_path / "paper"
    assert (
        main(
            [
                "-C",
                str(root),
                "annotate",
                "preview",
                "--only",
                "abstract_accuracy",
                "--out",
                "plan.json",
            ]
        )
        == 0
    )
    capsys.readouterr()
    with pytest.raises(SystemExit):
        main(["-C", str(root), "annotate", "create", "--plan", "plan.json", "--out", "copy.zip"])
    assert not (root / "copy.zip").exists()
    arguments = [
        "-C",
        str(root),
        "annotate",
        "create",
        "--plan",
        "plan.json",
        "--out",
        "copy.zip",
        "--attest-reviewed",
    ]
    assert main(arguments) == 0
    before = (root / "copy.zip").read_bytes()
    assert main(arguments) == 2
    assert (root / "copy.zip").read_bytes() == before


def test_forged_input_hashes_and_duplicate_selection_refused(tmp_path):
    create_demo(Project(tmp_path), "paper", document="docx")
    project = Project(tmp_path / "paper")
    with pytest.raises(PaperDeltaError):
        preview_annotations(project, {"occurrences": ["abstract_accuracy", "abstract_accuracy"]})
    plan = preview_annotations(project, {"occurrences": ["abstract_accuracy"]})
    forged = deepcopy(plan)
    forged["input_hashes"] = {}
    with pytest.raises(PaperDeltaError):
        annotation_bytes(project, forged)
